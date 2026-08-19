"""
FILE: build_phase5_current_prototype_doc.py
PURPOSE: Creates a beginner-friendly Word document that records SariSync's
current Phase 5 prototype architecture and capabilities.

This builder is only used to create the documentation file. It does not change
the SariSync website or its mock product data.
"""

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
OUTPUT = ROOT / "SariSync_Phase_5_Current_Prototype_Architecture.docx"

BLUE = "2E74B5"
DARK = "173E43"
MUTED = "5B6570"
LIGHT_BLUE = "E8EEF5"
LIGHT_GREEN = "E8F3EC"
LIGHT_ORANGE = "FFF1E6"
WHITE = "FFFFFF"
BLACK = "1E1E1E"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for col, width in zip(grid.gridCol_lst, widths):
        col.set(qn("w:w"), str(width))
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            set_cell_width(cell, width)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def set_keep_with_next(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    keep = OxmlElement("w:keepNext")
    p_pr.append(keep)


def configure_styles(doc):
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for style_name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, DARK, 10, 5),
    ):
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer.add_run("SariSync | Current Prototype Architecture")
    run.font.name = "Calibri"
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(MUTED)


def add_title(doc, text, subtitle):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run(text)
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    run.font.size = Pt(24)
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string(DARK)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(16)
    run = p.add_run(subtitle)
    run.font.name = "Calibri"
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor.from_string(MUTED)


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(text, style=f"Heading {level}")
    set_keep_with_next(p)
    return p


def add_body(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    if bold_prefix and text.startswith(bold_prefix):
        r = p.add_run(bold_prefix)
        r.bold = True
        p.add_run(text[len(bold_prefix):])
    else:
        p.add_run(text)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.25
        p.add_run(item)


def add_numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.25
        p.add_run(item)


def add_callout(doc, label, text, fill=LIGHT_ORANGE):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360])
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(label.upper())
    r.bold = True
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor.from_string(DARK)
    p2 = cell.add_paragraph(text)
    p2.paragraph_format.space_after = Pt(0)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_table(doc, headers, rows, widths, font_size=9.5):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_geometry(table, widths)
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for cell, label in zip(hdr.cells, headers):
        set_cell_shading(cell, LIGHT_BLUE)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(label)
        r.bold = True
        r.font.size = Pt(font_size)
        r.font.color.rgb = RGBColor.from_string(DARK)
    for row in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, row):
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(value)
            r.font.size = Pt(font_size)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_code_block(doc, text):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360])
    cell = table.cell(0, 0)
    set_cell_shading(cell, "F4F6F8")
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.05
    r = p.add_run(text)
    r.font.name = "Consolas"
    r._element.rPr.rFonts.set(qn("w:ascii"), "Consolas")
    r._element.rPr.rFonts.set(qn("w:hAnsi"), "Consolas")
    r.font.size = Pt(8.4)
    r.font.color.rgb = RGBColor.from_string(BLACK)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_page_break(doc):
    doc.add_page_break()


def build_document():
    doc = Document()
    configure_styles(doc)

    add_title(doc, "SariSync Phase 5", "Current Prototype Architecture and Capability Reference | August 2026")
    add_callout(
        doc,
        "Document purpose",
        "This is the updated reference for the SariSync browser prototype as it exists today. It replaces earlier target-only statements with the features, files, and limits currently present in the project.",
        LIGHT_GREEN,
    )
    add_heading(doc, "1. Current Position", 1)
    add_body(doc, "Phase 5 - Web Architecture is complete for the current prototype. The app has been reorganized so that pages, shared interface pieces, data helpers, storage helpers, and tests have clearer jobs.")
    add_body(doc, "The prototype remains a browser-based HTML, CSS, and JavaScript app. It uses sample product data and browser storage. It is not yet a live multi-device store system.")
    add_heading(doc, "What this document confirms", 2)
    add_bullets(doc, [
        "The current file structure and the purpose of each major file or folder.",
        "The screens and user actions that the prototype currently demonstrates.",
        "The features that are connected to mock data or local browser storage.",
        "The later work that still needs Supabase, real authentication, security, or device testing.",
    ])
    add_heading(doc, "2. Current Prototype Capability Summary", 1)
    add_table(doc, ["Area", "What works in the current prototype", "Data status"], [
        ("Price lookup", "Search products by partial name, category, or barcode; open product details.", "Mock product array"),
        ("Barcode controls", "Home and product forms offer scan controls; the camera scanner is added only when needed.", "Browser/device capability"),
        ("Product changes", "Add, edit, update retail or optional wholesale prices, and preserve current form behavior.", "Mock product array"),
        ("Price Pulse", "Shows recent price-change cards, summaries, empty state, and a route to Updates.", "Mock update data"),
        ("Browsing", "Browse products, filter by category, and open the standalone category view.", "Mock products and categories"),
        ("Settings", "Open profile, prepare CSV import/export actions, and show prototype setting feedback.", "Prototype controls"),
        ("Session display", "Reads the saved sariSyncSession and displays store/device names when available.", "localStorage"),
    ], [1800, 4900, 2660])
    add_callout(doc, "Important boundary", "CSV actions, barcode controls, session display, and navigation are connected in the prototype. They do not mean that a real Supabase store, secure account, cloud sync, or camera workflow has been production-tested.")

    add_page_break(doc)
    add_heading(doc, "3. Current Architecture Overview", 1)
    add_body(doc, "The prototype uses one shared browser shell for its main views. Page files coordinate each screen; shared files handle repeated interface parts or operations.")
    add_code_block(doc, "USER ACTION\n  -> Page controller (Home, Browse, Settings, etc.)\n  -> Shared component or service when needed\n  -> Product Repository / State / Local Storage\n  -> Mock data or saved browser session\n  -> Page re-renders the current result")
    add_heading(doc, "How to read the flow", 2)
    add_table(doc, ["Part", "Simple meaning", "Current example"], [
        ("Page", "A screen controller that listens for clicks and decides what to show.", "browse.js shows the product list and filters."),
        ("Component", "A reusable visual piece used by more than one screen.", "product-card.js creates product rows/cards."),
        ("Service", "A reusable operation that is not one specific screen.", "csv-service.js prepares import/export work."),
        ("Repository", "The middle file that reads or changes product data for pages.", "product-repository.js manages mock products."),
        ("State", "Temporary app memory for the current session and selected product.", "store.js remembers the selected product."),
        ("Storage", "A safe helper for browser localStorage.", "local-storage.js reads sariSyncSession."),
    ], [1700, 3500, 4160])
    add_heading(doc, "4. Current Design and Code Principles", 1)
    add_bullets(doc, [
        "One shared shell: index.html owns the main app space, one bottom navigation bar, and one floating Add Product button.",
        "One page, one main job: each file in src/js/pages focuses on a screen or workflow.",
        "Reuse repeated interface pieces: headers, product cards, product forms, Price Pulse, category preview, and toasts have dedicated component files.",
        "Keep product data behind Product Repository: pages ask it for products instead of directly editing the shared array.",
        "Use models to clean and check values: product and category helpers normalize names and validate inputs.",
        "Keep the UI responsive through shared CSS tokens, layout rules, and the responsive helper.",
    ])

    add_page_break(doc)
    add_heading(doc, "5. Actual Folder and File Structure", 1)
    add_body(doc, "This is the current application structure. A folder named README means the folder is intentionally explained or reserved for future files.")
    add_code_block(doc, "Digital Price Book/\n|-- index.html                         # Shared app shell and Home markup\n|-- settings.html                      # Redirect for old Settings links\n|-- pages/\n|   |-- auth/login.html                # Login and store-registration document\n|   `-- categories/category.html       # Standalone category document\n|-- src/\n|   |-- css/\n|   |   |-- styles.css                 # Loads CSS layers in order\n|   |   |-- tokens.css                 # Shared colors, spacing, radii, type\n|   |   |-- base.css                   # Basic browser and text rules\n|   |   |-- layout.css                 # Shell and page layout\n|   |   |-- components.css             # Components and responsive rules\n|   |   `-- pages/README.md            # Reserved for page-only CSS later\n|   |-- js/\n|   |   |-- navigation.js              # Tabs, active view, Back buttons\n|   |   |-- constants/app-constants.js # Repeated fixed values\n|   |   |-- state/store.js             # Session and selected product\n|   |   |-- storage/local-storage.js   # Browser storage helper\n|   |   |-- models/                    # Product and category validation\n|   |   |-- repositories/              # Product Repository\n|   |   |-- services/                  # Data, search, CSV, barcode, helpers\n|   |   `-- pages/                     # Individual page controllers\n|   |-- assets/README.md               # Local media location, if added later\n|   `-- README.md                      # Source-folder guide\n|-- tests/                             # Automated module checks\n`-- ARCHITECTURE.md                    # Short technical architecture reference")
    add_heading(doc, "6. Page Files and Their Current Responsibilities", 1)
    add_table(doc, ["Page file", "Current responsibility"], [
        ("home.js", "Coordinates the Home dashboard, search entry points, quick actions, Price Pulse, and routes to secondary views."),
        ("browse.js", "Shows browse/search/filter behavior for product browsing."),
        ("updates.js", "Shows the full list of recent price changes."),
        ("settings.js", "Shows settings controls and connects Import/Export actions."),
        ("profile.js", "Shows the current member profile screen."),
        ("product-details.js", "Shows lookup details, including retail and optional wholesale price."),
        ("add-product.js", "Opens the shared product form for a new product."),
        ("update-price.js", "Finds a product before opening the price/edit workflow."),
        ("edit-product.js", "Edits an existing product using the shared product form."),
        ("auth.js and category.js", "Support the separate login/registration and standalone category documents."),
    ], [2800, 6560])

    add_page_break(doc)
    add_heading(doc, "7. Shared Shell and Navigation", 1)
    add_body(doc, "index.html is the main app shell. It keeps the bottom navigation and floating add action in one place so Browse, Updates, and Settings do not each create their own navigation bar.")
    add_heading(doc, "Shell ownership", 2)
    add_bullets(doc, [
        "Home content and the reusable secondary-page area.",
        "One opaque bottom navigation bar for Home, Browse, Updates, and Settings.",
        "One floating Add Product action.",
        "The shared toast message area and CSV file input.",
        "Spacing that keeps scrollable content clear of the bottom navigation bar.",
    ])
    add_heading(doc, "Current navigation behavior", 2)
    add_table(doc, ["Destination", "How the user opens it", "Current return behavior"], [
        ("Home, Browse, Updates, Settings", "Bottom navigation", "Changes the active main view while keeping the shared shell."),
        ("Profile", "Home avatar or Settings option", "Current prototype returns to Settings."),
        ("Product Details", "Search, Browse, Updates, Price Pulse, or category list", "Returns to Home in the current implementation."),
        ("Add Product", "Floating Add button or Quick Action", "Back and a successful save return to Home."),
        ("Update Price", "Home Quick Action", "Back returns to Home; edit opens the Edit Item screen."),
        ("Edit Item", "Update Price, Browse, or Product Details", "Returns to the screen that opened the edit workflow."),
        ("Category Products", "Home category or Browse category", "Standalone category page; Back returns to Home."),
        ("Import / Export", "Settings controls", "Current actions run inside Settings rather than separate pages."),
    ], [2200, 3300, 3860])
    add_callout(doc, "Current limit", "The prototype uses simple route targets and page callbacks. It does not yet keep a complete browser-style Back history for every possible path.", LIGHT_ORANGE)
    add_heading(doc, "8. Shared Components", 1)
    add_table(doc, ["Component", "What it reuses"], [
        ("page-header.js", "The secondary-page label, title, description, and Back-button placement."),
        ("product-card.js", "Repeated product rows/cards used by listing and lookup screens."),
        ("product-form.js", "The shared Add Product and Edit Item fields, validation, barcode action, and save connection."),
        ("price-pulse.js", "Price Pulse summaries, sorting, change cards, and empty state on Home."),
        ("category-preview.js", "The compact category preview on Home."),
        ("toast.js", "Temporary success, warning, or error-style messages without redesigning the page."),
    ], [2700, 6660])

    add_page_break(doc)
    add_heading(doc, "9. Current State, Data, and Storage", 1)
    add_body(doc, "The current central state is intentionally small. It keeps information needed while the app is open; it does not yet act as a complete cloud data store.")
    add_table(doc, ["Current state item", "What it contains", "Where it comes from"], [
        ("session", "The saved sariSyncSession when it exists.", "localStorage through local-storage.js"),
        ("active store name", "The store name to display, or a safe fallback name.", "Session or app constants"),
        ("active device/user name", "The name shown for the current device/user, or a safe fallback.", "Session or app constants"),
        ("selected product", "The product most recently opened for details or editing.", "Temporary in-memory state"),
    ], [2200, 4200, 2960])
    add_heading(doc, "Product data flow", 2)
    add_code_block(doc, "Page or form\n  -> Product Repository\n  -> Product Model checks/normalizes the product\n  -> Mock products in services/data.js\n  -> Updated page display")
    add_body(doc, "The Product Repository currently works with mock products from data.js. This is deliberate: later, the repository can be changed to read Supabase without requiring every page to be rewritten.")
    add_heading(doc, "10. Models, Services, Repository, and Storage", 1)
    add_table(doc, ["Type", "Current files", "Current job"], [
        ("Models", "product.js, category.js", "Normalize names, optional prices, categories, and form values; validate required values."),
        ("Repository", "product-repository.js", "Get, find, add, and update mock products through one shared interface."),
        ("Services", "data.js, search-service.js, csv-service.js, barcode.js, responsive.js, utils.js", "Handle mock data, searching, CSV actions, barcode controls, responsive sizing, and formatting."),
        ("Storage", "local-storage.js", "Safely read, save, and remove browser localStorage values such as sariSyncSession."),
        ("Constants", "app-constants.js", "Keep repeated values such as device limit, CSV columns, defaults, and toast duration together."),
    ], [1700, 3550, 4110])

    add_page_break(doc)
    add_heading(doc, "11. Responsive UI and CSS", 1)
    add_body(doc, "The current prototype keeps its styling in shared CSS layers. The split makes it easier to find a design rule without changing unrelated parts of the app.")
    add_table(doc, ["CSS file", "Current responsibility"], [
        ("tokens.css", "Shared colors, spacing, border radius, and type values."),
        ("base.css", "Basic browser reset, typography, and accessibility-friendly defaults."),
        ("layout.css", "App shell, page container, navigation, and structural layout."),
        ("components.css", "Buttons, fields, cards, product forms, and responsive component rules."),
        ("styles.css", "Loads the CSS layers in the required order."),
    ], [2600, 6760])
    add_heading(doc, "Responsive behavior verified in the prototype", 2)
    add_bullets(doc, [
        "No horizontal overflow was found at 320 px, 768 px, or 1280 px widths during the Phase 5 review.",
        "The fixed bottom navigation bar remains opaque and page content has clearance above it.",
        "Cards, category previews, and Price Pulse respond to available width rather than relying only on one phone size.",
        "responsive.js provides a focused helper for viewport-derived sizing where the prototype uses it.",
    ])
    add_heading(doc, "12. Data, Authentication, and Security Boundary", 1)
    add_callout(doc, "Prototype versus production", "The current interface can demonstrate registration and login screens, but it does not provide production-ready account creation, authentication, Row Level Security, real store isolation, or multi-device synchronization.", LIGHT_ORANGE)
    add_table(doc, ["Current prototype uses", "Later production work needs"], [
        ("Mock products and categories in JavaScript", "Supabase tables, constraints, and store-specific records"),
        ("sariSyncSession in browser localStorage", "Supabase Auth session handling and secure profile records"),
        ("Existing Supabase client file left untouched", "Reviewed server-side security and Row Level Security policies"),
        ("Barcode controls in a browser", "Testing on supported physical devices with granted camera permission"),
        ("Prototype update display", "Realtime sync, offline queueing, and conflict decisions"),
    ], [4680, 4680])

    add_page_break(doc)
    add_heading(doc, "13. Verification Completed for Phase 5", 1)
    add_body(doc, "The following checks were completed during the final Phase 5 architecture review. They confirm the organization and prototype behavior, not production backend security.")
    add_table(doc, ["Check", "Result"], [
        ("JavaScript syntax", "Passed for all JavaScript files under src."),
        ("Automated module tests", "7 tests passed and 0 failed."),
        ("HTML, CSS, and JavaScript references", "Local references resolved to existing files."),
        ("Shared navigation", "Home, Browse, Updates, Settings, and secondary workflows were checked."),
        ("Back navigation", "Back buttons worked, including the context-aware Edit Item return path."),
        ("Search and category browsing", "Partial product search, Browse filters, and category search were checked."),
        ("Price Pulse and Updates", "Rendered from mock data with existing summaries and cards."),
        ("CSV actions", "Settings controls connect to the CSV service and preserve blank optional columns."),
        ("Barcode controls", "Home and product-form scan controls are present and connected."),
        ("Responsive layout", "No horizontal overflow found at 320 px, 768 px, and 1280 px."),
    ], [3100, 6260])
    add_heading(doc, "How to run the current automated tests", 2)
    add_numbered(doc, [
        "Open the Digital Price Book project in VS Code.",
        "Open the Terminal inside that project folder.",
        "Run: npm.cmd test",
        "Check that the terminal reports passing tests and no failures.",
    ])
    add_heading(doc, "14. Work Intentionally Deferred to Later Phases", 1)
    add_bullets(doc, [
        "Move mock product and category records into the approved Supabase database design.",
        "Build real email/password authentication and secure store membership access.",
        "Add Row Level Security so one store cannot read or change another store's data.",
        "Implement actual device registration, the five-device limit, and named updater records.",
        "Implement real synchronization, offline caching, queued edits, and conflict handling.",
        "Add an actual category management screen if the Version 1 scope is expanded to include it.",
        "Test the camera scanner, CSV file picker, and file downloads on physical supported devices.",
    ])

    add_page_break(doc)
    add_heading(doc, "15. Phase 5 Completion Checklist", 1)
    add_table(doc, ["Completion item", "Current status"], [
        ("Major pages separated into page files", "Completed"),
        ("Shared page/header/form/product card/Price Pulse/toast pieces separated", "Completed"),
        ("Product model and category model present", "Completed"),
        ("Product Repository protects direct product-array access", "Completed"),
        ("Search, CSV, barcode, formatting, and responsive helpers separated", "Completed"),
        ("Shared browser storage adapter and app state present", "Completed"),
        ("CSS layers separated and loaded in order", "Completed"),
        ("Beginner-friendly source-file comments present", "Completed"),
        ("Basic automated tests and test instructions available", "Completed"),
        ("Supabase production connection, authentication, and security", "Deferred to later phases"),
    ], [6500, 2860])
    add_heading(doc, "16. Phase 5 Handoff", 1)
    add_body(doc, "Phase 5 is complete for the SariSync prototype. The next implementation work can now use the Phase 6 approved data model and the separated architecture without changing the current interface simply to organize code.")
    add_callout(doc, "Recommended next phase", "Begin Phase 7 - Database Design only after the Phase 6 data-model document is marked approved. Phase 7 should convert approved entities and rules into real Supabase tables and constraints; it should not redesign the current prototype UI.", LIGHT_GREEN)
    add_heading(doc, "Appendix A. Beginner Glossary", 1)
    add_table(doc, ["Term", "Simple meaning"], [
        ("Prototype", "A working demonstration. It shows the app flow, but it is not yet the final real-data system."),
        ("Mock data", "Sample data stored in the code so screens can be tested without a database."),
        ("Component", "A reusable visual piece, such as a product card, form, or temporary toast message."),
        ("Service", "A helper for an operation, such as searching, exporting a CSV, or opening barcode controls."),
        ("Repository", "A data middle person. Pages ask it for products instead of reaching directly into a data array."),
        ("Model", "A helper that checks and cleans a record before it is saved or displayed."),
        ("State", "Temporary information the app remembers while it is open, such as the selected product."),
        ("localStorage", "Small browser storage used here for a prototype session; it is not a secure server database."),
        ("Supabase", "The planned backend platform for real data, authentication, security, and synchronization."),
        ("Toast", "A small temporary message, such as a save-success or error notice."),
    ], [2200, 7160])
    add_heading(doc, "Appendix B. Current Source Evidence", 1)
    add_bullets(doc, [
        "ARCHITECTURE.md: current file map, completed structure, verification results, and deferred work.",
        "src/js/state/store.js: current session and selected-product state responsibilities.",
        "src/js/navigation.js: current tab switching and Back-button behavior.",
        "tests/app-modules.test.js and tests/README.md: current automated module checks and run instructions.",
        "src/css and src/js folders: current modular source structure used by the browser prototype.",
    ])

    doc.core_properties.title = "SariSync Phase 5 Current Prototype Architecture"
    doc.core_properties.subject = "Current prototype architecture and capabilities"
    doc.core_properties.author = "SariSync project documentation"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
