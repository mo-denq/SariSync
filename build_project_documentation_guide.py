# FILE: build_project_documentation_guide.py | FOLDER: project root
# PURPOSE: Creates the beginner-friendly project documentation guide DOCX with
# phases, templates, checklists, examples, and diagram instructions.
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = 'SariSync_Project_Documentation_Guide.docx'
BLUE, DARK, LIGHT = '2E74B5', '1F4D78', 'E8EEF5'

def shade(cell, fill=LIGHT):
    tcPr = cell._tc.get_or_add_tcPr(); shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), fill); tcPr.append(shd)

def table(doc, headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers)); t.alignment = WD_TABLE_ALIGNMENT.LEFT; t.autofit = False
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]; c.width = Inches(widths[i] / 1440); shade(c)
        r = c.paragraphs[0].add_run(h); r.bold = True; r.font.color.rgb = RGBColor.from_string(DARK)
    for row in rows:
        cs = t.add_row().cells
        for i, value in enumerate(row):
            cs[i].width = Inches(widths[i] / 1440); cs[i].text = value; cs[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcPr = cs[i]._tc.get_or_add_tcPr(); tcW = OxmlElement('w:tcW'); tcW.set(qn('w:w'), str(widths[i])); tcW.set(qn('w:type'), 'dxa'); tcPr.append(tcW)
    doc.add_paragraph(); return t

def bullets(doc, values):
    for value in values: doc.add_paragraph(value, style='List Bullet')

def checklist(doc, values):
    for value in values: doc.add_paragraph('☐ ' + value)

def heading(doc, text, level=1): doc.add_heading(text, level)

doc = Document(); sec = doc.sections[0]
sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Inches(1)
sec.header_distance = sec.footer_distance = Inches(.492)
normal = doc.styles['Normal']; normal.font.name = 'Calibri'; normal.font.size = Pt(11); normal.paragraph_format.space_after = Pt(6); normal.paragraph_format.line_spacing = 1.25
for name, size, color, before, after in [('Heading 1',16,BLUE,18,10), ('Heading 2',13,BLUE,14,7), ('Heading 3',12,DARK,10,5)]:
    s = doc.styles[name]; s.font.name = 'Calibri'; s.font.size = Pt(size); s.font.bold = True; s.font.color.rgb = RGBColor.from_string(color); s.paragraph_format.space_before = Pt(before); s.paragraph_format.space_after = Pt(after)

p = doc.add_paragraph(); r = p.add_run('SariSync Project Documentation Guide'); r.bold = True; r.font.size = Pt(26); r.font.color.rgb = RGBColor.from_string(DARK)
p = doc.add_paragraph(); r = p.add_run('Example + fill-in template for planning, building, and maintaining the application'); r.font.size = Pt(15); r.font.color.rgb = RGBColor.from_string(BLUE)
doc.add_paragraph('This guide is written for the current HTML/CSS/JavaScript prototype. It also shows how the project can later evolve into a mobile application.')
table(doc, ['How to read this guide', 'Meaning'], [
    ('Example', 'Sample SariSync content that shows the expected level of detail.'),
    ('Your input', 'Replace the placeholder with your own decision, notes, link, or diagram.'),
    ('Deliverable', 'The file, diagram, decision, or working result produced by the phase.'),
    ('Complete when', 'A practical checklist for deciding whether the phase is finished.')], [2700, 6660])

heading(doc, 'How to use this document', 1)
bullets(doc, ['Read each phase before starting it.', 'Write your decisions in the Your input sections; do not leave important decisions only in chat messages.', 'Keep examples if they help, but clearly label which parts are still assumptions.', 'Link each diagram, spreadsheet, or design file from the relevant phase.', 'When a phase is complete, date it and record the document version.'])
heading(doc, 'Simple technical vocabulary', 1)
table(doc, ['Term', 'Simple meaning', 'SariSync example'], [
    ('View / page', 'A screen the user sees.', 'Home, Browse, Settings, or Profile.'),
    ('Component', 'A reusable interface piece.', 'A product card or settings tile.'),
    ('Markup', 'HTML structure that describes the screen.', 'A button containing a title and price.'),
    ('Render function', 'JavaScript that creates or updates markup.', 'renderProductCard(product).'),
    ('CSS class', 'A reusable styling label.', '.settings-card controls card appearance.'),
    ('Service', 'Logic that performs an operation.', 'Barcode scanning or CSV export.'),
    ('Repository', 'A layer that reads and writes data.', 'A product repository connected to Supabase.'),
    ('Data model', 'The structure of stored information.', 'Product fields such as retailPrice and category.'),
    ('API', 'A defined way for systems to communicate.', 'Supabase database requests.'),
    ('Acceptance criteria', 'Conditions proving a feature works.', 'A product saves without a wholesale price.')], [1900, 3900, 3560])

phases = [
('1. Project idea and problem definition', 'Explain the problem, users, setting, and value before deciding features.', ['What problem is costing users time or money?', 'Who experiences it most often?', 'What is the simplest useful solution?', 'How will success be recognized?'], ['Problem statement', 'Target user description', 'Goals and non-goals', 'Success measures'], 'Example: Rotating family members forget changing retail prices because wholesale costs change frequently. SariSync is a shared price lookup tool, not a full POS.', ['Problem statement is one clear paragraph.', 'Primary users and context are named.', 'Goals and non-goals are separated.', 'Success can be measured.']),
('2. Version 1 requirements and scope', 'Freeze what Version 1 includes, excludes, and limits.', ['Which features are essential for first release?', 'Which ideas belong in Version 2?', 'What limits keep the product focused?', 'What roles can perform each action?'], ['Feature scope', 'Roles', 'Limits', 'Functional and non-functional requirements'], 'Example: Version 1 includes search, barcode fallback, product price fields, Price Pulse, five devices, and offline cached lookup.', ['Version 1 document is approved.', 'Future ideas are separated.', 'Roles and limits are written.', 'Every requirement has a testable outcome.']),
('3. User-flow planning', 'Map how users move from starting action to completed result.', ['Where does each flow start?', 'What does the user tap or type?', 'What does the system do?', 'Where can the flow succeed, fail, or stop?'], ['Lookup, update, add, profile, import, export, and navigation flows'], 'Example: Home → Search Canton → Results → Product Lookup → Back to Home.', ['Every major flow has entry and exit points.', 'Back behavior is defined.', 'Success, error, and empty states are included.', 'No flow ends at a dead end.']),
('4. Screen and information architecture', 'Decide which screens exist and what information belongs on each one.', ['What is the purpose of each screen?', 'What belongs in the shared shell?', 'Which screens are separate views?', 'What is the priority of information?'], ['Screen inventory', 'Navigation map', 'Screen hierarchy', 'Content priority notes'], 'Example: Home prioritizes search and Price Pulse; Settings contains Store Management and Profile access.', ['Every view has one main purpose.', 'Primary actions are visible.', 'Shared navigation is documented.', 'No screen contains unrelated administration.']),
('5. UI/UX and design system', 'Define how the app looks, responds, and stays consistent.', ['What are the colors, spacing, radii, and type scale?', 'How large must prices be?', 'How does the layout adapt to phones and tablets?', 'How are loading, empty, and error states shown?'], ['Design tokens', 'Wireframes', 'High-fidelity screens', 'Component inventory'], 'Example: Use a high-contrast price scale, flexible card widths, and one reusable Settings Tile style.', ['A design reference exists.', 'Responsive rules are written.', 'Reusable components are identified.', 'Accessibility states are included.']),
('6. Data modeling', 'Define the shape of every important record before database work.', ['What fields does a Product need?', 'Which values are required or optional?', 'How is a price change recorded?', 'What identifies a store, member, and device?'], ['Product, Category, PriceUpdate, Profile, Store, and UserRole models'], 'Example Product fields: barcode, name, category, costPrice, retailPrice, wholesalePrice?, bottleDeposit?, notes, updatedBy, updatedAt.', ['Required and optional fields are marked.', 'Relationships are documented.', 'Example records exist.', 'Invalid values and empty values are defined.']),
('7. Technical architecture', 'Show how the browser/app, services, database, and device features connect.', ['Which layer owns UI?', 'Which layer owns business logic?', 'Where is data stored?', 'What happens when the network is unavailable?'], ['Architecture diagram', 'Layer responsibilities', 'Technology decisions'], 'Example: HTML/CSS/JS view → product service → repository → localStorage cache or Supabase.', ['Dependencies are identified.', 'Each layer has a responsibility.', 'Data flow is diagrammed.', 'Temporary prototype choices are labeled.']),
('8. File and project structure', 'Make it easy to find and safely change code.', ['Which files represent pages?', 'Which code is reusable?', 'Where do services and models live?', 'What naming rules will be followed?'], ['Folder tree', 'File ownership notes', 'Naming conventions'], 'Example: pages/home/home.js; services/barcode.js; models/product.js; components/settings-tile.js.', ['Folders have clear purposes.', 'Files have one primary responsibility.', 'Shared code is not duplicated.', 'A new developer can find the main files.']),
('9. Backend and Supabase planning', 'Define production data, authentication, storage, and security.', ['Which tables are needed?', 'How are stores isolated?', 'What policies protect data?', 'Which files or images use Storage?'], ['Schema, policies, API/service notes, environment setup'], 'Example tables: stores, profiles, devices, products, price_updates; every product is scoped to a storeId.', ['Tables and relationships are approved.', 'Access policies are planned.', 'Secrets are not placed in public code.', 'Prototype data migration is considered.']),
('10. Offline and synchronization planning', 'Define how the app behaves without internet and resolves later sync.', ['What data is cached?', 'How are offline edits queued?', 'What status is shown?', 'What if two devices edit the same product?'], ['Offline flow', 'Queue rules', 'Conflict rule', 'Sync status states'], 'Example: Save locally → mark pending → retry after connection → broadcast update → show updatedBy and updatedAt.', ['Offline lookup works by design.', 'Queue and retry behavior is written.', 'Conflict behavior is chosen.', 'Failure messaging is clear.']),
('11. Implementation', 'Turn the approved requirements and designs into working code.', ['Which feature should be built first?', 'What can be reused?', 'What is the smallest testable increment?', 'What must not be changed?'], ['Feature branches/commits', 'Working screens', 'Code comments', 'Implementation notes'], 'Example order: session → product model → Home search → add/edit → Price Pulse → sync.', ['Each feature meets its acceptance criteria.', 'Console errors are resolved.', 'Shared styles remain consistent.', 'Implementation notes explain intentional deviations.']),
('12. Testing and quality assurance', 'Check behavior, layout, data, accessibility, and resilience.', ['What should happen for valid and invalid inputs?', 'Does it work at different widths?', 'Does it work offline?', 'Can a user recover from errors?'], ['Test plan', 'Test cases', 'Bug list', 'Regression checklist'], 'Example: Test empty wholesale price, denied camera permission, duplicate barcode, offline save, and five-device limit.', ['Critical flows pass.', 'Responsive screens are checked.', 'Known bugs are recorded.', 'A release candidate is repeatable.']),
('13. User testing', 'Observe real family members using the app in the store context.', ['Can a parent find a price without help?', 'Do labels make sense?', 'Are prices readable in glare or low light?', 'Where do users hesitate?'], ['Test script', 'Participant notes', 'Issue priority list', 'Before/after decisions'], 'Example task: “Find the retail price for Lucky Me Canton and tell me who last changed it.”', ['Real users complete core tasks.', 'Observed issues are prioritized.', 'Changes are linked to evidence.', 'Retesting confirms improvements.']),
('14. Deployment and maintenance', 'Define how the app is released, monitored, backed up, and supported.', ['Where is the production app hosted?', 'How are database changes deployed?', 'How are backups made?', 'Who responds to incidents?'], ['Deployment checklist', 'Environment notes', 'Backup plan', 'Maintenance log'], 'Example: Deploy static web assets, apply Supabase migrations, verify auth, test sync, and keep a rollback copy.', ['Release steps are repeatable.', 'Backups are tested.', 'Environment secrets are documented safely.', 'An owner for maintenance is assigned.']),
('15. Future versions and roadmap', 'Keep new ideas visible without destabilizing Version 1.', ['What user problem is next?', 'What evidence supports it?', 'Does it require new data or permissions?', 'What moves from idea to committed scope?'], ['Version 2 backlog', 'Prioritized roadmap', 'Decision log'], 'Example future ideas: supplier price history, advanced roles, analytics, multi-store accounts, and richer import templates.', ['Version 2 ideas are separated from V1.', 'Priorities have reasons.', 'Dependencies are noted.', 'A review date is scheduled.'])
]

for title, purpose, questions, inputs, example, done in phases:
    heading(doc, title, 1)
    doc.add_paragraph('Purpose: ' + purpose)
    heading(doc, 'Questions to answer', 2); bullets(doc, questions)
    heading(doc, 'Your input', 2); bullets(doc, ['Write your decision here: ________________________________', 'Add a link, drawing, spreadsheet, or note here: ________________________________'] + inputs)
    heading(doc, 'Example — SariSync', 2); doc.add_paragraph(example)
    heading(doc, 'Deliverables', 2); bullets(doc, inputs)
    heading(doc, 'Complete when', 2); checklist(doc, done)

heading(doc, 'Documentation templates', 1)
doc.add_paragraph('Copy these headings into separate project files or keep them as sections in one master document.')
templates = [
('Project overview', 'Your input: Product name, one-sentence description, current technology, audience, owner, version, last updated.'),
('Goals and non-goals', 'Your input: What the product must accomplish; what it intentionally will not do.'),
('Target users', 'Your input: User type, environment, skill level, needs, frustrations, and success measure.'),
('User roles', 'Your input: Role name, permissions, restrictions, and examples of who has the role.'),
('Feature scope', 'Your input: Feature name, purpose, priority, owner, status, and Version 1/Version 2 label.'),
('System limits', 'Your input: Device, file, product, photo, performance, storage, and permission limits.'),
('Functional requirements', 'Your input: “The system shall…” statements with acceptance criteria.'),
('Non-functional requirements', 'Your input: Performance, privacy, accessibility, reliability, portability, and security expectations.'),
('User flows', 'Your input: Entry, action, response, next view, success, error, empty state, and exit.'),
('Screen inventory', 'Your input: Screen name, purpose, entry points, exit points, primary action, and states.'),
('Data models', 'Your input: Field name, type, required/optional, example, validation, and relationship.'),
('Navigation structure', 'Your input: Shared shell, tabs, pages, back behavior, and deep links.'),
('File documentation', 'Your input: File path, responsibility, dependencies, public functions, and change notes.'),
('API and service documentation', 'Your input: Service name, inputs, outputs, errors, authentication, and data source.'),
('Testing checklist', 'Your input: Test ID, setup, action, expected result, actual result, status, and bug link.'),
('Known issues', 'Your input: Issue, impact, workaround, priority, owner, and planned fix.'),
('Future improvements', 'Your input: Idea, user problem, evidence, dependencies, priority, and target version.'),
('Version history', 'Your input: Version, date, changes, author, known issues, and approval.')
]
for name, text in templates:
    heading(doc, name, 2); doc.add_paragraph(text); doc.add_paragraph('Example: ______________________________________________')

heading(doc, 'Diagram guide', 1)
doc.add_paragraph('A diagram is a visual explanation. You can draw it by hand, use Figma/FigJam, or use Mermaid text. Use rectangles for views, arrows for movement, diamonds for decisions, and notes for errors.')
diagrams = [
('User-flow diagram', 'Shows the steps a user follows.', 'Flowchart', 'Home → Search Canton → Results → Product Lookup → Back.', 'Your diagram: [start] → [action] → [response] → [success/error]'),
('Navigation map', 'Shows how views connect.', 'Site map / flowchart', 'Home connects to Browse, Updates, Settings, Profile, and product flows.', 'Your diagram: [shell] → [tab] → [view]'),
('Screen hierarchy', 'Shows parent/child relationships between views.', 'Tree diagram', 'App Shell → Home; App Shell → Browse; App Shell → Settings → Profile.', 'Your diagram: [root] → [child] → [subview]'),
('System architecture', 'Shows UI, services, storage, and external systems.', 'Layered architecture', 'Browser UI → Product Service → Repository → localStorage/Supabase.', 'Your diagram: [client] → [service] → [data layer]'),
('Data relationship diagram', 'Shows records and their relationships.', 'ERD', 'Store has many Devices and Products; Product has many Price Updates.', 'Your diagram: [entity] 1—many [entity]'),
('Offline synchronization flow', 'Shows what happens when connection disappears and returns.', 'State/flow diagram', 'Edit offline → Local cache → Pending queue → Reconnect → Supabase → Other devices.', 'Your diagram: [offline] → [queue] → [retry] → [synced]'),
('Authentication flow', 'Shows sign-in, registration, joining, and session states.', 'Sequence or flowchart', 'Register → Verify mobile/PIN → Create store → Home; Join → Store code → Device nickname → Home.', 'Your diagram: [credential] → [verification] → [session]'),
('Product price update flow', 'Shows how a saved price becomes an audit record.', 'Sequence diagram', 'Search product → Edit retail price → Save → Update previous/current → Save updatedBy/updatedAt → Price Pulse.', 'Your diagram: [user] → [form] → [service] → [database] → [feed]')
]
for name, purpose, kind, example, blank in diagrams:
    heading(doc, name, 2)
    table(doc, ['Item', 'Guide'], [('Purpose', purpose), ('Recommended type', kind), ('Shapes/content', 'Rectangles = views or entities; arrows = movement; diamonds = decisions; notes = errors.'), ('Example', example), ('Your blank template', blank)], [2100, 7260])

heading(doc, 'Final documentation checklist', 1)
checklist(doc, ['Project overview and Version 1 scope are approved.', 'Every major screen has a purpose, entry, exit, and states.', 'The main user flows are diagrammed and written.', 'Models, services, repositories, and file responsibilities are documented.', 'Security, offline, sync, and data-portability decisions are recorded.', 'Testing, deployment, maintenance, and future-version notes exist.', 'Each document has an owner and last-updated date.'])

doc.save(OUT)
print(OUT)
