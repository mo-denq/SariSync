from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


REFERENCE = Path(
    r"C:\Users\Yasmien\.codex\plugins\cache\openai-curated-remote\openai-templates\0.1.1\skills\artifact-template-legal-memorandum\assets\reference.docx"
)
OUTPUT_DIR = Path(
    r"C:\Users\Yasmien\OneDrive - sdca.edu.ph\Documents\Digital Price Book\.codex-artifacts\client-outreach\outputs"
)

BLACK = RGBColor(0, 0, 0)
GRAY = RGBColor(118, 118, 118)
LINK_BLUE = "305496"


def set_cell_shading(paragraph, fill="F4F4F4"):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = p_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        p_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_bottom_rule(paragraph, color="666666", size="8"):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "5")
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)


def set_left_rule(paragraph, color="B7B7B7", size="10"):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), size)
    left.set(qn("w:space"), "8")
    left.set(qn("w:color"), color)
    p_bdr.append(left)


def set_run_font(run, size=11, bold=False, italic=False, color=BLACK):
    run.font.name = "Garamond"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Garamond")


def remove_body_content(document):
    body = document._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def configure_document(document, header_text):
    remove_body_content(document)
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.5)
    section.footer_distance = Inches(0.5)
    section.different_first_page_header_footer = False

    normal = document.styles["Normal"]
    normal.font.name = "Garamond"
    normal.font.size = Pt(11)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Garamond")

    if "Template Block" not in [style.name for style in document.styles]:
        block_style = document.styles.add_style("Template Block", WD_STYLE_TYPE.PARAGRAPH)
        block_style.font.name = "Garamond"
        block_style.font.size = Pt(10.5)
        block_style._element.rPr.rFonts.set(qn("w:eastAsia"), "Garamond")
        block_style.paragraph_format.left_indent = Inches(0.22)
        block_style.paragraph_format.right_indent = Inches(0.12)
        block_style.paragraph_format.space_after = Pt(0)
        block_style.paragraph_format.line_spacing = 1.0

    header = section.header
    while len(header.paragraphs) > 1:
        p = header.paragraphs[-1]._element
        p.getparent().remove(p)
    hp = header.paragraphs[0]
    hp.clear()
    hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    hr = hp.add_run(header_text.upper())
    set_run_font(hr, size=9, bold=False, color=GRAY)


def add_title(document, title, size=22):
    p = document.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(16)
    r = p.add_run(title.upper())
    set_run_font(r, size=size, bold=True)


def add_metadata(document, label, value, last=False):
    p = document.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.0
    label_run = p.add_run(f"{label}: ")
    set_run_font(label_run, size=10.5, bold=True)
    value_run = p.add_run(value)
    set_run_font(value_run, size=10.5)
    if last:
        p.paragraph_format.space_after = Pt(12)
        set_bottom_rule(p)
    return p


def add_section_heading(document, text, *, page_break_before=False):
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.page_break_before = page_break_before
    r = p.add_run(text.upper())
    set_run_font(r, size=12, bold=True)
    r.font.underline = True
    return p


def add_body(document, text, *, italic=False, shaded=False, keep=False):
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.08
    p.paragraph_format.keep_together = keep
    r = p.add_run(text)
    set_run_font(r, size=11, italic=italic)
    if shaded:
        set_cell_shading(p)
        set_left_rule(p)
        p.paragraph_format.left_indent = Inches(0.16)
        p.paragraph_format.right_indent = Inches(0.10)
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(8)
    return p


def add_label_body(document, label, text):
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.08
    p.paragraph_format.keep_together = True
    r1 = p.add_run(f"{label}: ")
    set_run_font(r1, size=11, bold=True)
    r2 = p.add_run(text)
    set_run_font(r2, size=11)
    return p


def add_bullet(document, text, *, level=0):
    p = document.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.28 + (0.22 * level))
    p.paragraph_format.first_line_indent = Inches(-0.18)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.04
    p.paragraph_format.keep_together = True
    bullet = p.add_run("•\u00a0 ")
    set_run_font(bullet, size=11, bold=True)
    r = p.add_run(text)
    set_run_font(r, size=11)
    return p


def add_numbered(document, number, text):
    p = document.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.30)
    p.paragraph_format.first_line_indent = Inches(-0.24)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.04
    p.paragraph_format.keep_together = True
    lead = p.add_run(f"{number}. ")
    set_run_font(lead, size=11, bold=True)
    rest = p.add_run(text)
    set_run_font(rest, size=11)
    return p


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    rel_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    new_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), LINK_BLUE)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_fonts = OxmlElement("w:rFonts")
    r_fonts.set(qn("w:ascii"), "Garamond")
    r_fonts.set(qn("w:hAnsi"), "Garamond")
    r_pr.extend([r_fonts, color, underline])
    new_run.append(r_pr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    new_run.append(text_node)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)
    return hyperlink


def add_link_line(document, label, url):
    p = document.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.28)
    p.paragraph_format.first_line_indent = Inches(-0.18)
    p.paragraph_format.space_after = Pt(4)
    bullet = p.add_run("•\u00a0 ")
    set_run_font(bullet, size=11, bold=True)
    label_run = p.add_run(f"{label}: ")
    set_run_font(label_run, size=11, bold=True)
    add_hyperlink(p, url, url)
    return p


def add_template_block(document, lines):
    for index, line in enumerate(lines):
        p = document.add_paragraph(style="Template Block")
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.keep_together = True
        p.paragraph_format.keep_with_next = index < len(lines) - 1
        if index == 0:
            p.paragraph_format.space_before = Pt(3)
        if index == len(lines) - 1:
            p.paragraph_format.space_after = Pt(8)
        set_cell_shading(p, "F4F4F4")
        set_left_rule(p)
        r = p.add_run(line if line else " ")
        set_run_font(r, size=10.5)


def build_services_reference():
    document = Document(REFERENCE)
    configure_document(document, "Internal Reference • Client Outreach")
    add_title(document, "Developer Services and Sales Reference", size=21)
    add_metadata(document, "Prepared for", "Client outreach coordinator")
    add_metadata(document, "Development partner", "CJ Sacdalan")
    add_metadata(document, "Updated", "September 6, 2026")
    add_metadata(
        document,
        "Purpose",
        "Approved working information for discussing possible client projects",
        last=True,
    )

    add_body(
        document,
        "This is the main internal reference for outreach. Use it to keep messages accurate and consistent. It is not a client contract or legal advice. If a request is not covered here, record it and ask the developer before making a commitment.",
        shaded=True,
    )

    add_section_heading(document, "1. Service Focus")
    add_body(
        document,
        "The development partner can consider a broad range of custom websites and software based on the client's actual needs. The focus is software—not large hardware projects. Every request still needs a feasibility review before it is accepted.",
    )
    add_bullet(document, "Business websites and web-based systems")
    add_bullet(document, "Workflow and task automation")
    add_bullet(document, "Analytics, dashboards, reporting, and business monitoring")
    add_bullet(document, "AI-assisted software features when appropriate")
    add_bullet(document, "Other custom software that matches the developer's skills and available capacity")

    add_section_heading(document, "2. Suitable Clients")
    add_body(
        document,
        "There is no single required industry. A good prospect is a business with a clear problem that software could realistically improve. Do not tell a prospect that every request is automatically possible; first understand the problem and let the developer review the proposed solution.",
    )
    add_bullet(document, "A manual or repetitive workflow is taking too much time")
    add_bullet(document, "Important information or customer data is difficult to monitor")
    add_bullet(document, "Customers struggle to find information, contact the business, book, order, or make an inquiry")
    add_bullet(document, "The business needs clearer reporting to understand what is working")

    add_section_heading(document, "3. Scope and Pricing")
    add_body(
        document,
        "There are currently no fixed packages or service tiers. The client explains the problem, requested features, users, and expected result first. The developer then reviews the scope before giving a price.",
    )
    add_bullet(document, "Do not lowball or offer a price only to win the client.")
    add_bullet(document, "Normally, do not include a price in the first outreach email.")
    add_bullet(document, "You may discuss a client's budget range, but do not present it as the final project price.")
    add_bullet(document, "Any discount, payment plan, or fixed quote needs developer approval.")

    add_section_heading(document, "4. Benefits You May Discuss")
    add_body(document, "Describe possible benefits carefully and connect them to the client's real situation:")
    add_bullet(document, "Simpler workflows and fewer manual steps")
    add_bullet(document, "Easier monitoring of operations, customer activity, or other useful data")
    add_bullet(document, "Better visibility into what is working and what may need improvement")
    add_bullet(document, "Information that can support marketing decisions or customer targeting")
    add_bullet(document, "A solution configured around the business's actual process")
    add_body(
        document,
        "Never guarantee more sales, a specific return, perfect automation, or another result that has not been proven and approved.",
        shaded=True,
    )

    add_section_heading(document, "5. Timeline and Deposit")
    add_label_body(document, "Working timeline", "A basic project is expected to take at least one month. This is not a promised completion date. The developer approves the actual schedule after reviewing the scope and workload.")
    add_label_body(document, "Deposit", "The developer indicated either a 30% or 50% down payment. The rule for choosing between them is still unconfirmed, so ask before stating the exact deposit to a client.")

    add_section_heading(document, "6. Portfolio Material")
    add_link_line(document, "Nextstep — Digital products and business systems", "https://nextstep.nxtstudio.workers.dev")
    add_link_line(document, "CJ Sacdalan — Full-stack mobile application developer", "https://cjsacdalan.com")
    add_body(
        document,
        "These links may be shared as portfolio material. They are not automatically testimonials or proof of a specific business result. Do not invent clients, performance numbers, testimonials, or outcomes.",
    )

    add_section_heading(document, "7. Your Role: Client Outreach Coordinator")
    add_bullet(document, "Research possible clients using public business information.")
    add_bullet(document, "Write and manually send personalized outreach messages.")
    add_bullet(document, "Listen to the client's current problem and ask simple discovery questions.")
    add_bullet(document, "Suggest a high-level direction based on the stated problem, without guaranteeing feasibility.")
    add_bullet(document, "Track replies, follow-ups, requests, and opt-outs.")
    add_bullet(document, "Prepare a clear handoff summary for the developer.")
    add_body(
        document,
        "Your job is to start and manage the conversation—not to act as the developer or make technical and commercial promises on the developer's behalf.",
        shaded=True,
    )

    add_section_heading(document, "8. Developer and Shared Responsibilities")
    add_label_body(document, "Developer", "Checks technical feasibility, confirms the solution, sets the final price and timeline, and performs the development work.")
    add_label_body(document, "Shared", "Agree on the written scope, deliverables, responsibilities, payment terms, and how changes will be handled before the project starts.")

    add_section_heading(document, "9. Items You Must Confirm Before Promising", page_break_before=True)
    add_bullet(document, "Final price, discount, or payment plan")
    add_bullet(document, "Final deadline, start date, or rush delivery")
    add_bullet(document, "Whether a requested feature or integration is technically possible")
    add_bullet(document, "Hosting, maintenance, support, ownership, or ongoing fees")
    add_bullet(document, "Guaranteed results, performance, or business outcomes")
    add_bullet(document, "Any feature or deliverable that is not in the approved written scope")

    add_section_heading(document, "10. Questions to Ask an Interested Prospect")
    questions = [
        "What is the main problem you want to solve?",
        "How do you handle that process now, and which part takes the most time?",
        "What result would make the project useful for your business?",
        "Which features do you already know you need?",
        "Who will use the website or system?",
        "Do you have a preferred launch date?",
        "Do you have a budget range in mind? You may say that a custom quote still requires a scope review.",
    ]
    for i, question in enumerate(questions, 1):
        add_numbered(document, i, question)

    add_section_heading(document, "11. Scope Changes")
    add_body(
        document,
        "List all agreed features and deliverables in the project contract or written scope. If the client later asks for something new, record the request and have the developer review its price and timeline before agreeing. This protects both the client and the team from uncontrolled scope growth.",
    )

    add_section_heading(document, "12. Open Items to Confirm With the Developer")
    add_bullet(document, "Minimum acceptable project budget or a rule for identifying lowball offers")
    add_bullet(document, "When the deposit is 30% and when it is 50%")
    add_bullet(document, "Who sends the proposal and contract, receives payment, and issues receipts or invoices")
    add_bullet(document, "Your commission, referral fee, or share and when it is paid")
    add_bullet(document, "What must be received before work starts, such as the deposit, signed scope, content, or account access")

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    output = OUTPUT_DIR / "Developer Services and Sales Reference.docx"
    document.save(output)
    return output


def build_email_guide():
    document = Document(REFERENCE)
    configure_document(document, "Internal Communication Guide • Client Outreach")
    add_title(document, "Email Templates and Communication Rules", size=20)
    add_metadata(document, "Prepared for", "Client outreach coordinator")
    add_metadata(document, "Development partner", "CJ Sacdalan")
    add_metadata(document, "Updated", "September 6, 2026")
    add_metadata(document, "Purpose", "Safe, accurate templates for manual client communication", last=True)

    add_body(
        document,
        "These templates are starting points. Replace every bracketed field, remove lines that do not apply, and check every statement before sending. Keep outreach manual and personalized.",
        shaded=True,
    )

    add_section_heading(document, "1. Before You Email")
    add_bullet(document, "Confirm that the business appears active using a current public source.")
    add_bullet(document, "Write down one real observation or possible opportunity. Do not invent a problem.")
    add_bullet(document, "Use only a clearly published business email or official contact form. Never guess an email address.")
    add_bullet(document, "Check the tracker to make sure the business has not opted out or already been contacted.")
    add_bullet(document, "Do not scrape LinkedIn, automate LinkedIn messages, or send messages automatically.")

    add_section_heading(document, "2. First Outreach Email")
    add_template_block(
        document,
        [
            "Subject: A possible way to improve [specific process] at [Business name]",
            "",
            "Hi [Name or Business team],",
            "",
            "I came across [Business name] through [public source] and noticed [one accurate observation].",
            "",
            "I work with CJ Sacdalan, a software developer who builds custom websites and business systems, including automation, analytics, and AI-assisted tools. Based on [observation], one possible improvement could be [simple idea] to help [realistic benefit].",
            "",
            "Would you be open to a brief email conversation about how you currently handle [process or problem]? If this is not relevant, please let me know and I will not follow up.",
            "",
            "Regards,",
            "[Your name]",
            "Client outreach coordinator working with CJ Sacdalan",
            "[Your business contact information]",
        ],
    )

    add_section_heading(document, "3. When a Prospect Is Interested")
    add_template_block(
        document,
        [
            "Subject: Re: [Original subject]",
            "",
            "Hi [Name],",
            "",
            "Thank you for replying. To understand the request before I involve the developer, could you tell me:",
            "1. What is the main problem you want to solve?",
            "2. How do you handle it now, and which part takes the most time?",
            "3. What result would make the project useful for you?",
            "4. Are there features you already know you need, and who will use them?",
            "5. Do you have a preferred launch date or budget range?",
            "",
            "A custom quote and delivery schedule will be confirmed after the developer reviews the scope.",
            "",
            "Regards,",
            "[Your name]",
        ],
    )

    add_section_heading(document, "4. When You Need a Technical Answer")
    add_template_block(
        document,
        [
            "Thanks for explaining that. I do not want to give you an inaccurate answer, so I will confirm the technical details with our developer and reply after the request has been reviewed.",
        ],
    )

    add_section_heading(document, "5. When the Prospect Asks for a Price")
    add_template_block(
        document,
        [
            "Our work is priced according to the actual features and scope rather than a fixed package. If you share the main problem, required features, intended users, and preferred timing, I can prepare the details for a custom quote from the developer.",
        ],
    )

    add_section_heading(document, "6. One Polite Follow-Up")
    add_template_block(
        document,
        [
            "Subject: Following up — [short topic]",
            "",
            "Hi [Name or Business team],",
            "",
            "I wanted to follow up on my note about [specific idea]. If improving [process] is not a priority right now, no problem—I will close my note and will not continue following up.",
            "",
            "Regards,",
            "[Your name]",
        ],
    )

    add_section_heading(document, "7. No-Interest or Opt-Out Reply")
    add_template_block(
        document,
        [
            "Thank you for letting me know. I have noted your preference and will not contact you again about this. Wishing you and the team well.",
        ],
    )

    add_section_heading(document, "8. Handoff Summary for the Developer")
    add_template_block(
        document,
        [
            "Business: [Business name, industry, and location]",
            "Contact: [Name, role, and public business contact used]",
            "Current problem: [Client's own words or a careful summary]",
            "Current process: [How they handle it now]",
            "Desired result: [What they want to improve]",
            "Requested features: [List; mark uncertain requests clearly]",
            "Users: [Who will use the system]",
            "Preferred timing: [Date or 'not provided']",
            "Budget range: [Range or 'not provided']",
            "Promises already made: [Ideally 'none']",
            "Questions for developer: [Feasibility, price, timeline, technical questions]",
            "Source links and notes: [URLs and useful context]",
        ],
    )

    add_section_heading(document, "9. Communication Rules")
    add_bullet(document, "Tell the truth about the business, the portfolio, and what has been approved.")
    add_bullet(document, "Do not claim testimonials, results, or clients that cannot be documented.")
    add_bullet(document, "Do not finalize price, discounts, deadlines, or technical feasibility without developer approval.")
    add_bullet(document, "Do not promise that a suggested feature will definitely solve the problem.")
    add_bullet(document, "Personalize each message; do not send the same generic message to a large list.")
    add_bullet(document, "Send no more than one polite follow-up unless the prospect continues the conversation.")
    add_bullet(document, "Stop immediately when a business declines or asks not to be contacted, and mark it in the tracker.")
    add_bullet(document, "Record every sent message, reply, follow-up date, and handoff in the lead tracker.")

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    output = OUTPUT_DIR / "Email Templates and Communication Rules.docx"
    document.save(output)
    return output


if __name__ == "__main__":
    built = [build_services_reference(), build_email_guide()]
    for path in built:
        print(path)
