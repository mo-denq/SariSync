from pathlib import Path
from zipfile import ZipFile

from docx import Document


OUTPUT_DIR = Path(
    r"C:\Users\Yasmien\OneDrive - sdca.edu.ph\Documents\Digital Price Book\.codex-artifacts\client-outreach\outputs"
)


for path in sorted(OUTPUT_DIR.glob("*.docx")):
    document = Document(path)
    text = "\n".join(p.text for p in document.paragraphs)
    lower = text.lower()
    with ZipFile(path) as package:
        document_xml = package.read("word/document.xml").decode("utf-8", errors="ignore")
        footer_parts = [
            package.read(name).decode("utf-8", errors="ignore")
            for name in package.namelist()
            if name.startswith("word/footer") and name.endswith(".xml")
        ]
        header_parts = [
            package.read(name).decode("utf-8", errors="ignore")
            for name in package.namelist()
            if name.startswith("word/header") and name.endswith(".xml")
        ]
    print(
        {
            "file": path.name,
            "paragraphs": len(document.paragraphs),
            "sections": len(document.sections),
            "words": len(text.split()),
            "contains_lorem": "lorem ipsum" in lower,
            "contains_old_legal_title": "legal memo" in lower,
            "has_page_field": any("PAGE" in part for part in footer_parts),
            "has_internal_header": any("Internal" in part or "INTERNAL" in part for part in header_parts),
            "hyperlinks": document_xml.count("w:hyperlink"),
        }
    )
