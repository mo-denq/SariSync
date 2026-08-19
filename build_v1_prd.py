# FILE: build_v1_prd.py | FOLDER: project root
# PURPOSE: Builds the Version 1 requirements document covering features, roles,
# limits, functional requirements, and non-functional requirements.
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = 'SariSync_V1_Feature_Roles_Requirements.docx'
BLUE, DARK, FILL = '2E74B5', '1F4D78', 'E8EEF5'

def shade(cell, fill=FILL):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), fill); tcPr.append(shd)

def table(doc, headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers)); t.alignment = WD_TABLE_ALIGNMENT.LEFT; t.autofit = False
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]; c.width = Inches(widths[i] / 1440); shade(c)
        r = c.paragraphs[0].add_run(h); r.bold = True; r.font.color.rgb = RGBColor.from_string(DARK)
    for row in rows:
        cells = t.add_row().cells
        for i, value in enumerate(row):
            cells[i].width = Inches(widths[i] / 1440); cells[i].text = value
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcPr = cells[i]._tc.get_or_add_tcPr()
            tcW = OxmlElement('w:tcW'); tcW.set(qn('w:w'), str(widths[i])); tcW.set(qn('w:type'), 'dxa'); tcPr.append(tcW)
    doc.add_paragraph()
    return t

def bullet(doc, text):
    doc.add_paragraph(text, style='List Bullet')

def number(doc, text):
    doc.add_paragraph(text, style='List Number')

doc = Document()
section = doc.sections[0]
section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Inches(1)
section.header_distance = section.footer_distance = Inches(.492)
normal = doc.styles['Normal']; normal.font.name = 'Calibri'; normal.font.size = Pt(11)
normal.paragraph_format.space_after = Pt(6); normal.paragraph_format.line_spacing = 1.25
for name, size, color, before, after in [('Heading 1',16,BLUE,18,10), ('Heading 2',13,BLUE,14,7), ('Heading 3',12,DARK,10,5)]:
    s = doc.styles[name]; s.font.name = 'Calibri'; s.font.size = Pt(size); s.font.bold = True; s.font.color.rgb = RGBColor.from_string(color)
    s.paragraph_format.space_before = Pt(before); s.paragraph_format.space_after = Pt(after)

p = doc.add_paragraph(); r = p.add_run('SariSync Version 1'); r.bold = True; r.font.size = Pt(26); r.font.color.rgb = RGBColor.from_string(DARK)
p = doc.add_paragraph(); r = p.add_run('Feature Scope, Roles, Limits, and Requirements'); r.font.size = Pt(15); r.font.color.rgb = RGBColor.from_string(BLUE)
doc.add_paragraph('Product requirements draft for the shared digital price book for family-run sari-sari stores.')
table(doc, ['Document field', 'Version 1 decision'], [
    ('Product', 'SariSync shared digital price book'),
    ('Audience', 'Family members rotating shifts in one sari-sari store'),
    ('Release intent', 'Fast price lookup and shared updates, not a full POS'),
    ('Status', 'Draft for product, design, and implementation alignment')], [2700, 6660])

doc.add_heading('1. Product direction', 1)
doc.add_paragraph('SariSync helps family members answer one question quickly: “How much is this item right now?” Version 1 prioritizes fast lookup, shared updates, clear freshness information, and simple onboarding. Sales logging, accounting, and complex inventory operations remain outside the first release.')
doc.add_heading('Goals', 2)
for x in ['Reduce incorrect prices caused by changing wholesale costs.', 'Give authorized family devices one shared price book.', 'Make lookup possible in a few taps, including offline lookup.', 'Show who changed a price and when.']: bullet(doc, x)
doc.add_heading('Non-goals', 2)
for x in ['No transaction ledger, cash drawer, profit-and-loss reporting, or full POS checkout.', 'No supplier ordering, accounting, loyalty program, or customer storefront.', 'No mandatory barcode requirement.', 'No advanced permission matrix beyond owner and store member in Version 1.']: bullet(doc, x)

doc.add_heading('2. Version 1 feature scope', 1)
table(doc, ['Area', 'Capability', 'Priority'], [
    ('Authentication', 'Owner registration and member join using mobile/store code and PIN/password.', 'Must'),
    ('Shared access', 'Up to five devices per store; each device has a nickname for audit display.', 'Must'),
    ('Home', 'Greeting, sync status, search, barcode scan, Price Pulse, quick actions, category preview.', 'Must'),
    ('Products', 'Barcode, name, category, cost, retail, optional wholesale, bottle deposit, notes, photo.', 'Must'),
    ('Browse/update', 'Browse categories, search products, edit product details and prices.', 'Must'),
    ('Profile', 'Display name, mobile/email details, profile photo placeholder, device nickname.', 'Should'),
    ('Import/export', 'CSV import preparation and CSV backup under Settings.', 'Should'),
    ('Offline/realtime', 'Cached lookup, queued edits, and synchronized product changes.', 'Must/Should')], [2200, 5700, 1460])

doc.add_heading('3. Roles and responsibilities', 1)
table(doc, ['Role', 'Can do', 'Limits'], [
    ('Account owner', 'Create store; manage identity, access, products, profile, import, and export.', 'Cannot access another store without its code/invitation.'),
    ('Store member', 'Join store; search, scan, browse, and edit products when shared editing is enabled.', 'Cannot change account ownership or core store identity.'),
    ('Device identity', 'Choose a nickname such as Mom or Caleb shown on saved updates.', 'Not a separate permission role or user account.')], [1800, 4800, 2760])
doc.add_paragraph('Recommended default: store members may update prices because the app supports rotating family shifts. Owner-only controls should focus on access, store identity, and portability.')

doc.add_heading('4. Version 1 limits and guardrails', 1)
table(doc, ['Limit', 'Rule', 'Reason'], [
    ('Devices', 'Maximum 5 active devices per store.', 'Fits the small family-operated store.'),
    ('Store scope', 'One account represents one store location.', 'Prevents mixed price books.'),
    ('Photos', 'One optional photo per product.', 'Supports visual verification without a media library.'),
    ('Wholesale', 'Optional; blank values display as Not set.', 'Not every item is sold or bought in bulk.'),
    ('Barcode', 'Optional; scan or type when available.', 'Local/repackaged items may have none.'),
    ('Categories', 'Use the nine approved store categories.', 'Keeps browsing consistent.'),
    ('Audit', 'Store updatedBy and updatedAt for each edit.', 'Provides transparency without a full ledger.'),
    ('Import', 'CSV first; spreadsheet templates later.', 'Reduces first-release complexity.')], [1800, 4500, 3060])

doc.add_heading('5. Functional requirements', 1)
doc.add_heading('Account and access', 2)
for x in ['The owner can register full name, required mobile number, optional email, store name, location, optional exact address/photo, and PIN/password.', 'A family member can join with a store code and choose a device nickname.', 'A sixth active device is blocked until an existing device is removed.', 'The session persists across refreshes until sign-out.']: bullet(doc, x)
doc.add_heading('Products and pricing', 2)
for x in ['Item name, category, and retail price are required.', 'Cost, wholesale, bottle deposit, notes, barcode, and photo are optional where stated.', 'Names are normalized consistently; the current prototype saves names in uppercase.', 'Price edits retain current price, updater, timestamp, and previous price when available.', 'The interface never invents a ₱0.00 old price.']: bullet(doc, x)
doc.add_heading('Search, browse, and scan', 2)
for x in ['Search supports partial product names and may match category or barcode.', 'Barcode scanning uses the camera when supported and manual entry as fallback.', 'Browse supports category filters and a dedicated edit flow.', 'Current retail price is visually dominant.']: bullet(doc, x)
doc.add_heading('Sync and offline', 2)
for x in ['The app shows synced, syncing, or waiting-for-connection status.', 'Cached data remains readable offline.', 'Offline edits are queued and retried after reconnection.', 'Before production, document conflict handling; recommended initial rule is latest valid update wins while retaining updater and timestamp.']: bullet(doc, x)

doc.add_heading('6. Non-functional requirements', 1)
for x in ['Usability: reach a current retail price from Home in a few taps.', 'Accessibility: high contrast, readable prices, labels beyond color, and touch targets around 44px or larger.', 'Responsive layout: support small phones, large phones, and tablets with flexible widths and a safe max-width.', 'Performance: avoid rebuilding unrelated screens when one product changes.', 'Privacy/security: authenticated store sessions and no cross-store data exposure.', 'Reliability: cached lookup works during temporary network loss.', 'Portability: CSV export uses the agreed product columns and preserves blank optional fields.']: bullet(doc, x)

doc.add_heading('7. Version 1 acceptance criteria', 1)
for x in ['Owner registration reaches Home.', 'Devices two through five can join and identify themselves.', 'Searching Canton finds matching variants with current retail prices.', 'Supported barcode scans open a matching product or show a clear no-match message.', 'A product saves successfully without a wholesale price.', 'A price edit shows new price, updater, time, and valid old price when available.', 'Price Pulse shows recent changes without a false old price.', 'Cached lookup remains available offline and queued edits sync later.', 'Import/export is available from Settings, not the daily Home dashboard.', 'Names, prices, forms, and navigation do not clip on supported screens.']: number(doc, x)

doc.add_heading('8. Recommended release sequence', 1)
for x in ['Stabilize account/session and local product storage.', 'Finalize product schema, price history fields, and category values.', 'Complete Home search, barcode, Browse, Add Product, and Update Price.', 'Add offline queue and conflict behavior.', 'Connect realtime sync and validate five-device behavior.', 'Validate CSV import/export using the store’s real spreadsheet.', 'Run family-user testing at the counter in bright and low-light conditions.']: number(doc, x)

doc.add_heading('9. Open decisions before production', 1)
for x in ['Can store members always edit, or can the owner disable editing?', 'What conflict rule applies when two devices edit offline?', 'What exact CSV column names and date format are permanent?', 'How will profile photos be stored and limited?', 'What is account recovery if the owner loses the PIN and has no email?']: bullet(doc, x)

doc.save(OUT)
print(OUT)
