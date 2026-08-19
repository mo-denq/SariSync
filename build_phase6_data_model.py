# FILE: build_phase6_data_model.py | FOLDER: project root
# PURPOSE: Generates the Phase 6 data-model documentation, including entities,
# fields, relationships, and implementation notes for the SariSync prototype.
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parent
OUT = ROOT / "SariSync_Phase_6_Data_Model.docx"
DIAGRAM = ROOT / "work" / "phase6_relationship_diagram.png"

NAVY = "173C3B"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
ORANGE = "F25A24"
PALE_ORANGE = "FFF0E8"
PALE_BLUE = "E8EEF5"
PALE_GREEN = "E8F5EC"
PALE_GOLD = "FFF6D8"
LIGHT = "F7F8F6"
GRID = "D7DDD9"
MUTED = "667673"
WHITE = "FFFFFF"
BLACK = "172C2B"
RED = "A13D32"

PAGE_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=120, bottom=90, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_table_geometry(table, widths):
    assert sum(widths) == PAGE_WIDTH_DXA
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(PAGE_WIDTH_DXA))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT_DXA))
    tbl_ind.set(qn("w:type"), "dxa")
    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths[idx]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)


def set_run(run, size=11, bold=False, color=BLACK, font="Calibri", italic=False):
    run.font.name = font
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), font)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), font)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    set_run(run, size=9, color=MUTED)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    paragraph._p.append(fld)


def keep_with_next(paragraph):
    paragraph.paragraph_format.keep_with_next = True


def set_cell_text(cell, text, bold=False, color=BLACK, size=9.2, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.12
    r = p.add_run(str(text))
    set_run(r, size=size, bold=bold, color=color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(doc, headers, rows, widths, header_fill=PALE_BLUE, size=9.2):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for i, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], header, bold=True, color=NAVY, size=9.2)
        set_cell_shading(table.rows[0].cells[i], header_fill)
    set_repeat_table_header(table.rows[0])
    for row_data in rows:
        row = table.add_row()
        for i, value in enumerate(row_data):
            align = WD_ALIGN_PARAGRAPH.CENTER if headers[i] in {"Required?", "Status", "Type"} else WD_ALIGN_PARAGRAPH.LEFT
            set_cell_text(row.cells[i], value, size=size, align=align)
    set_table_geometry(table, widths)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    return table


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(text, style=f"Heading {level}")
    keep_with_next(p)
    return p


def add_body(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.25
    if bold_lead and text.startswith(bold_lead):
        r1 = p.add_run(bold_lead)
        set_run(r1, bold=True)
        r2 = p.add_run(text[len(bold_lead):])
        set_run(r2)
    else:
        r = p.add_run(text)
        set_run(r)
    return p


def add_callout(doc, label, text, fill=PALE_ORANGE, accent=ORANGE):
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    set_table_geometry(table, [PAGE_WIDTH_DXA])
    set_cell_shading(table.cell(0, 0), fill)
    cell = table.cell(0, 0)
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(label.upper())
    set_run(r, size=9, bold=True, color=accent)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    p2.paragraph_format.line_spacing = 1.2
    r2 = p2.add_run(text)
    set_run(r2, size=10.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_simple_list(doc, items, numbered=False):
    style = "List Number" if numbered else "List Bullet"
    for item in items:
        p = doc.add_paragraph(style=style)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.25
        r = p.add_run(item)
        set_run(r)


def create_diagram():
    DIAGRAM.parent.mkdir(parents=True, exist_ok=True)
    w, h = 1700, 1050
    img = Image.new("RGB", (w, h), "white")
    draw = ImageDraw.Draw(img)
    try:
        title_font = ImageFont.truetype("arialbd.ttf", 44)
        box_font = ImageFont.truetype("arialbd.ttf", 31)
        sub_font = ImageFont.truetype("arial.ttf", 22)
        label_font = ImageFont.truetype("arial.ttf", 20)
    except OSError:
        title_font = box_font = sub_font = label_font = ImageFont.load_default()

    draw.text((70, 35), "SariSync Version 1 - Logical Relationship Model", fill=f"#{NAVY}", font=title_font)

    boxes = {
        "auth.users": (90, 135, 390, 255, PALE_GOLD, "Supabase-managed login"),
        "profiles": (520, 135, 820, 255, PALE_ORANGE, "One record per user"),
        "store_members": (945, 135, 1305, 255, PALE_BLUE, "Connects users and stores"),
        "stores": (1370, 135, 1630, 255, PALE_GREEN, "Shared price book"),
        "devices": (190, 430, 490, 550, LIGHT, "Phone, tablet or browser"),
        "categories": (600, 430, 900, 550, LIGHT, "Product grouping"),
        "products": (1010, 430, 1310, 550, PALE_ORANGE, "Current product details"),
        "price_updates": (745, 745, 1105, 865, PALE_GREEN, "Price history and audit"),
    }

    def box(name):
        x1, y1, x2, y2, fill, subtitle = boxes[name]
        draw.rounded_rectangle((x1, y1, x2, y2), radius=22, fill=f"#{fill}", outline=f"#{GRID}", width=3)
        tw = draw.textbbox((0, 0), name, font=box_font)[2]
        draw.text(((x1 + x2 - tw) / 2, y1 + 24), name, fill=f"#{NAVY}", font=box_font)
        sw = draw.textbbox((0, 0), subtitle, font=sub_font)[2]
        draw.text(((x1 + x2 - sw) / 2, y1 + 75), subtitle, fill=f"#{MUTED}", font=sub_font)

    for name in boxes:
        box(name)

    def arrow(a, b, label, a_side="right", b_side="left"):
        ax1, ay1, ax2, ay2, _, _ = boxes[a]
        bx1, by1, bx2, by2, _, _ = boxes[b]
        points = {
            "left": lambda x1, y1, x2, y2: (x1, (y1 + y2) // 2),
            "right": lambda x1, y1, x2, y2: (x2, (y1 + y2) // 2),
            "top": lambda x1, y1, x2, y2: ((x1 + x2) // 2, y1),
            "bottom": lambda x1, y1, x2, y2: ((x1 + x2) // 2, y2),
        }
        start = points[a_side](ax1, ay1, ax2, ay2)
        end = points[b_side](bx1, by1, bx2, by2)
        draw.line((start, end), fill=f"#{MUTED}", width=4)
        ex, ey = end
        if b_side == "left":
            arrow_head = [(ex, ey), (ex - 14, ey - 9), (ex - 14, ey + 9)]
        elif b_side == "right":
            arrow_head = [(ex, ey), (ex + 14, ey - 9), (ex + 14, ey + 9)]
        elif b_side == "top":
            arrow_head = [(ex, ey), (ex - 9, ey - 14), (ex + 9, ey - 14)]
        else:
            arrow_head = [(ex, ey), (ex - 9, ey + 14), (ex + 9, ey + 14)]
        draw.polygon(arrow_head, fill=f"#{MUTED}")
        mx, my = (start[0] + end[0]) // 2, (start[1] + end[1]) // 2
        bbox = draw.textbbox((0, 0), label, font=label_font)
        lw = bbox[2] - bbox[0]
        draw.rectangle((mx - lw / 2 - 7, my - 15, mx + lw / 2 + 7, my + 15), fill="white")
        draw.text((mx - lw / 2, my - 12), label, fill=f"#{MUTED}", font=label_font)

    arrow("auth.users", "profiles", "1 to 1")
    arrow("profiles", "store_members", "1 to many")
    arrow("store_members", "stores", "many to 1")
    arrow("profiles", "devices", "1 to many", "bottom", "top")
    arrow("stores", "categories", "1 to many", "bottom", "top")
    arrow("stores", "products", "1 to many", "bottom", "top")
    arrow("categories", "products", "1 to many")
    arrow("products", "price_updates", "1 to many", "bottom", "top")
    arrow("profiles", "price_updates", "records updater", "bottom", "left")
    arrow("devices", "price_updates", "records device", "bottom", "left")

    draw.text((70, 960), "Reading guide: 1 to many means one record on the first side can connect to several records on the second side.", fill=f"#{MUTED}", font=label_font)
    img.save(DIAGRAM, dpi=(180, 180))


def configure_styles(doc):
    sec = doc.sections[0]
    sec.page_width = Inches(8.5)
    sec.page_height = Inches(11)
    sec.top_margin = Inches(1)
    sec.bottom_margin = Inches(1)
    sec.left_margin = Inches(1)
    sec.right_margin = Inches(1)
    sec.header_distance = Inches(0.492)
    sec.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(BLACK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, DARK_BLUE, 10, 5),
    ):
        style = doc.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for list_name in ("List Bullet", "List Number"):
        style = doc.styles[list_name]
        style.font.name = "Calibri"
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(.375)
        style.paragraph_format.first_line_indent = Inches(-.188)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.25

    header = sec.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = hp.add_run("SariSync | Phase 6 Data Model")
    set_run(r, size=9, bold=True, color=MUTED)
    footer = sec.footer
    add_page_number(footer.paragraphs[0])


def add_cover(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(110)
    p.paragraph_format.space_after = Pt(16)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("PHASE 6")
    set_run(r, size=11, bold=True, color=ORANGE)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run("SariSync Version 1\nData Model")
    set_run(r, size=30, bold=True, color=NAVY)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(44)
    r = p.add_run("A beginner-friendly blueprint for organizing SariSync data before database implementation")
    set_run(r, size=13, color=MUTED, italic=True)

    add_callout(
        doc,
        "Document status",
        "Draft for review. Existing Supabase work: stores and products tables, UUID primary keys, and products.store_id referencing stores.id. Recommended entities and rules in this document must be reviewed before Phase 7.",
        fill=PALE_ORANGE,
    )

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(80)
    r = p.add_run("Prepared for the SariSync shared digital price book project")
    set_run(r, size=10, color=MUTED)
    doc.add_page_break()


def entity_tables(doc):
    entities = [
        ("Profile", "A person who uses SariSync.", "profiles", "Recommended"),
        ("Store", "One family-operated store and its shared price book.", "stores", "Exists"),
        ("Store Member", "Connects a user to a store and records their role.", "store_members", "Recommended"),
        ("Device", "A phone, tablet, or browser used by a family member.", "devices", "Recommended"),
        ("Category", "A product grouping used by one store.", "categories", "Recommended"),
        ("Product", "The current description and prices of one store item.", "products", "Exists"),
        ("Price Update", "A history record showing an old price, new price, updater, and time.", "price_updates", "Recommended"),
    ]
    add_table(doc, ["Entity", "Simple purpose", "Suggested table", "Status"], entities, [1550, 4450, 1900, 1460])


DATA = {
    "Stores": [
        ("id", "UUID", "Yes", "Unique store identifier", "Primary key"),
        ("store_code", "Text", "Yes", "Code used to recognize or join a store", "Unique"),
        ("store_name", "Text", "Yes", "Store name displayed in the app", ""),
        ("barangay", "Text", "Yes", "Store barangay", ""),
        ("city", "Text", "Yes", "City or municipality", ""),
        ("province", "Text", "Yes", "Province", ""),
        ("street_address", "Text", "No", "Optional exact address", "Nullable"),
        ("photo_url", "Text", "No", "Optional store-photo location", "Nullable"),
        ("created_at", "Timestamp", "Yes", "Creation date and time", "Default: current time"),
        ("updated_at", "Timestamp", "Yes", "Latest store-information change", "Default: current time"),
    ],
    "Profiles": [
        ("id", "UUID", "Yes", "Authenticated user identifier", "Primary key; later references auth.users.id"),
        ("full_name", "Text", "Yes", "Name shown to family members", ""),
        ("mobile_number", "Text", "Yes", "Login, verification, or contact number", "Final login method decided in Phase 8"),
        ("email", "Text", "No", "Optional recovery email", "Nullable"),
        ("photo_url", "Text", "No", "Optional profile-photo location", "Nullable"),
        ("created_at", "Timestamp", "Yes", "Profile creation time", "Default: current time"),
        ("updated_at", "Timestamp", "Yes", "Latest profile change", "Default: current time"),
    ],
    "Store Members": [
        ("id", "UUID", "Yes", "Membership identifier", "Primary key"),
        ("store_id", "UUID", "Yes", "Store being accessed", "References stores.id"),
        ("user_id", "UUID", "Yes", "User who belongs to the store", "References profiles.id"),
        ("role", "Text", "Yes", "Owner or member", "Controlled list"),
        ("status", "Text", "Yes", "Active, invited, or removed", "Controlled list"),
        ("joined_at", "Timestamp", "Yes", "Membership start time", "Default: current time"),
    ],
    "Devices": [
        ("id", "UUID", "Yes", "Device identifier", "Primary key"),
        ("store_id", "UUID", "Yes", "Store using the device", "References stores.id"),
        ("user_id", "UUID", "Yes", "User using the device", "References profiles.id"),
        ("nickname", "Text", "Yes", "Name shown in update history", "Example: Mom's Phone"),
        ("last_active_at", "Timestamp", "No", "Latest device activity", "Nullable"),
        ("is_active", "Boolean", "Yes", "Whether device access is allowed", "Default: true"),
        ("created_at", "Timestamp", "Yes", "Device registration time", "Default: current time"),
    ],
    "Categories": [
        ("id", "UUID", "Yes", "Category identifier", "Primary key"),
        ("store_id", "UUID", "Yes", "Store that owns the category", "References stores.id"),
        ("name", "Text", "Yes", "Displayed category name", "Unique within one store"),
        ("display_order", "Integer", "Yes", "Position in category lists", "Default value may be used"),
        ("is_active", "Boolean", "Yes", "Whether category can be selected", "Default: true"),
        ("created_at", "Timestamp", "Yes", "Creation time", "Default: current time"),
    ],
    "Products": [
        ("id", "UUID", "Yes", "Product identifier", "Primary key"),
        ("store_id", "UUID", "Yes", "Store that owns the product", "References stores.id; not unique"),
        ("category_id", "UUID", "Yes", "Selected category", "References categories.id"),
        ("barcode", "Text", "No", "Optional product barcode", "Nullable; unique within a store when present"),
        ("item_name", "Text", "Yes", "Product name", "Displayed in uppercase"),
        ("cost_price", "Numeric", "No", "Store purchase cost", "Nullable; cannot be negative"),
        ("retail_price", "Numeric", "Yes", "Current selling price", "Cannot be negative"),
        ("wholesale_price", "Numeric", "No", "Optional bulk price", "Nullable; cannot be negative"),
        ("bottle_deposit", "Numeric", "No", "Optional container deposit", "Nullable; cannot be negative"),
        ("notes", "Text", "No", "Additional product details", "Nullable"),
        ("photo_url", "Text", "No", "Optional product-photo location", "Nullable"),
        ("created_by", "UUID", "Yes", "User who added the product", "References profiles.id"),
        ("updated_by", "UUID", "Yes", "User who last edited the product", "References profiles.id"),
        ("created_at", "Timestamp", "Yes", "Creation time", "Default: current time"),
        ("updated_at", "Timestamp", "Yes", "Latest edit time", "Default: current time"),
        ("is_active", "Boolean", "Yes", "Active or archived product", "Default: true"),
    ],
    "Price Updates": [
        ("id", "UUID", "Yes", "Price-update identifier", "Primary key"),
        ("store_id", "UUID", "Yes", "Store where change occurred", "References stores.id"),
        ("product_id", "UUID", "Yes", "Product that changed", "References products.id"),
        ("previous_retail_price", "Numeric", "No", "Retail price before change", "Nullable for a new product"),
        ("new_retail_price", "Numeric", "Yes", "New retail price", "Cannot be negative"),
        ("previous_wholesale_price", "Numeric", "No", "Previous wholesale price", "Nullable"),
        ("new_wholesale_price", "Numeric", "No", "New wholesale price", "Nullable"),
        ("updated_by_user_id", "UUID", "Yes", "User who made the change", "References profiles.id"),
        ("device_id", "UUID", "No", "Device used for the change", "References devices.id; nullable"),
        ("notes", "Text", "No", "Optional explanation", "Nullable"),
        ("created_at", "Timestamp", "Yes", "Update time", "Default: current time"),
    ],
}


def build():
    create_diagram()
    doc = Document()
    configure_styles(doc)
    add_cover(doc)

    add_heading(doc, "1. How to Use This Document", 1)
    add_body(doc, "Phase 6 defines what information SariSync needs, what each piece means, and how different pieces connect. Phase 7 will convert this approved plan into Supabase tables, columns, constraints, and indexes.")
    add_callout(doc, "Important distinction", "This document is the data blueprint. It is not a SQL script and does not change the live Supabase project.")

    add_heading(doc, "1.1 Essential Technical Terms", 2)
    terms = [
        ("Entity", "A type of thing SariSync stores, such as a Product or Store."),
        ("Table", "The database implementation of an entity."),
        ("Attribute / field", "One detail about an entity, such as retail price."),
        ("Column", "The database implementation of an attribute."),
        ("Row / record", "One saved example, such as one Lucky Me product."),
        ("Primary key", "A unique identifier for one row."),
        ("Foreign key", "A column that connects one table to another."),
        ("Relationship", "The connection between entities."),
        ("Cardinality", "How many records can participate, such as one store having many products."),
        ("Nullable", "The field is optional and may contain no value."),
        ("Constraint", "A database-enforced rule, such as preventing a negative price."),
        ("UUID", "A long, automatically generated unique identifier."),
        ("Timestamp", "A stored date and time."),
        ("Schema", "The overall structure of tables, columns, keys, and relationships."),
        ("Data dictionary", "A reference explaining every field in the model."),
    ]
    add_table(doc, ["Term", "Simple meaning"], terms, [2100, 7260], header_fill=PALE_BLUE, size=9.7)

    add_heading(doc, "2. Current Phase 6 Status", 1)
    add_body(doc, "Based on the current public schema, SariSync already has the two central tables and their core relationship.")
    status = [
        ("stores table", "Completed", "Contains store identity, name, owner contact, and location fields."),
        ("products table", "Completed", "Contains the required product and price fields."),
        ("UUID primary keys", "Completed", "Both tables use UUID identifiers."),
        ("products.store_id foreign key", "Completed", "References public.stores.id."),
        ("Additional Version 1 entities", "Not started", "Profiles, store members, devices, categories, and price updates still need approval."),
        ("Plain-text PIN approach", "Needs revision", "Do not treat the current stores.pin design as production-ready."),
    ]
    add_table(doc, ["Item", "Status", "Meaning"], status, [2500, 1700, 5160], header_fill=PALE_GREEN)
    add_callout(doc, "Current schema warning", "Verify that products.store_id is not unique. One store must be allowed to own many product rows. Keep the existing foreign key to stores.id.", fill=PALE_GOLD, accent="7A5A00")

    add_heading(doc, "3. Entity List", 1)
    add_body(doc, "An entity is a type of thing the application needs to remember. The table name is the suggested Phase 7 implementation.")
    entity_tables(doc)

    add_heading(doc, "4. Data Dictionary", 1)
    add_body(doc, "The following tables describe the recommended Version 1 attributes. Status: recommended means the design must be reviewed before it is implemented.")
    for name, rows in DATA.items():
        add_heading(doc, name, 2)
        purpose = {
            "Stores": "Represents one family-operated sari-sari store and its shared price book.",
            "Profiles": "Stores the in-app information of one authenticated person.",
            "Store Members": "Connects people to stores and records their roles.",
            "Devices": "Identifies the phone, tablet, or browser used to make changes.",
            "Categories": "Maintains consistent product groupings for one store.",
            "Products": "Stores current product descriptions and prices.",
            "Price Updates": "Preserves price history for Price Pulse and the Updates screen.",
        }[name]
        add_body(doc, purpose)
        add_table(doc, ["Field", "Type", "Required?", "Meaning", "Rule / relationship"], rows, [1700, 1200, 1000, 2780, 2680], size=8.4)

    add_heading(doc, "5. Logical Relationship Diagram", 1)
    add_body(doc, "This diagram shows the proposed connections. It is a logical model: it explains the business structure before exact SQL implementation.")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run()
    diagram_shape = r.add_picture(str(DIAGRAM), width=Inches(6.45))
    diagram_shape._inline.docPr.set(
        "descr",
        "SariSync logical relationship diagram connecting authentication users, profiles, store members, stores, devices, categories, products, and price updates.",
    )
    p.paragraph_format.space_after = Pt(6)
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(8)
    rr = cap.add_run("Figure 1. Recommended SariSync Version 1 entity relationships.")
    set_run(rr, size=9, italic=True, color=MUTED)

    add_heading(doc, "6. Relationship List", 1)
    relationships = [
        ("Authentication user - Profile", "One to one", "Each authenticated Supabase user has one SariSync profile."),
        ("Profile - Store Member", "One to many", "One person may have one or more store memberships."),
        ("Store - Store Member", "One to many", "One store may contain several family members."),
        ("Profile - Device", "One to many", "One person may use more than one device."),
        ("Store - Device", "One to many", "One store may allow up to five active devices in Version 1."),
        ("Store - Category", "One to many", "Each store manages its own category list."),
        ("Store - Product", "One to many", "One store owns many products; one product belongs to one store."),
        ("Category - Product", "One to many", "One category contains many products."),
        ("Product - Price Update", "One to many", "One product may have a complete history of price changes."),
        ("Profile - Price Update", "One to many", "A user may create many price updates."),
        ("Device - Price Update", "One to many", "A device may be recorded on many updates."),
    ]
    add_table(doc, ["Relationship", "Cardinality", "Simple explanation"], relationships, [2900, 1700, 4760], size=9)

    add_heading(doc, "7. Primary Key Plan", 1)
    add_body(doc, "A primary key is the unique ID of a row. UUID is recommended because SariSync will synchronize data created on different devices.")
    primary = [
        ("profiles", "id", "UUID", "Same identity used by the authenticated user."),
        ("stores", "id", "UUID", "Unique store identity."),
        ("store_members", "id", "UUID", "Unique membership identity."),
        ("devices", "id", "UUID", "Unique registered-device identity."),
        ("categories", "id", "UUID", "Unique category identity."),
        ("products", "id", "UUID", "Unique product identity."),
        ("price_updates", "id", "UUID", "Unique price-history identity."),
    ]
    add_table(doc, ["Table", "Primary key", "Type", "Purpose"], primary, [2000, 1800, 1100, 4460])

    add_heading(doc, "8. Foreign Key Plan", 1)
    add_body(doc, "A foreign key connects a row to an existing row in another table. These links will later support both data accuracy and Row Level Security.")
    fks = [
        ("profiles.id", "auth.users.id", "Connect profile to authenticated login"),
        ("store_members.store_id", "stores.id", "Identify the joined store"),
        ("store_members.user_id", "profiles.id", "Identify the member"),
        ("devices.store_id", "stores.id", "Identify the device's store"),
        ("devices.user_id", "profiles.id", "Identify the device user"),
        ("categories.store_id", "stores.id", "Keep categories separated by store"),
        ("products.store_id", "stores.id", "Keep products separated by store"),
        ("products.category_id", "categories.id", "Assign a valid category"),
        ("products.created_by", "profiles.id", "Record who added the product"),
        ("products.updated_by", "profiles.id", "Record who last edited it"),
        ("price_updates.store_id", "stores.id", "Identify the store"),
        ("price_updates.product_id", "products.id", "Identify the changed product"),
        ("price_updates.updated_by_user_id", "profiles.id", "Identify the updater"),
        ("price_updates.device_id", "devices.id", "Identify the device when available"),
    ]
    add_table(doc, ["Foreign-key column", "References", "Reason"], fks, [3100, 2600, 3660], size=9)

    add_heading(doc, "9. Required and Optional Rules", 1)
    add_body(doc, "Required means the value must be supplied. Optional means the database may store NULL, which means no value was provided.")
    rules = [
        ("Product item name", "Required", "A product cannot be recognized without a name."),
        ("Product category", "Required", "Every active product must belong to a valid category."),
        ("Retail price", "Required", "This is the app's primary lookup value."),
        ("Barcode", "Optional", "Some local or repacked products have no barcode."),
        ("Cost price", "Optional", "May be unavailable during onboarding."),
        ("Wholesale price", "Optional", "Not every product is sold in bulk."),
        ("Bottle deposit", "Optional", "Only some containers require a deposit."),
        ("Notes and photos", "Optional", "Helpful but not required for price lookup."),
        ("Previous price", "Optional", "A newly added product has no previous price."),
        ("Store membership", "Required for access", "A user must be an active member before accessing store data."),
        ("Device nickname", "Required", "Used to explain which device made an update."),
    ]
    add_table(doc, ["Field or rule", "Status", "Reason"], rules, [2550, 1900, 4910], size=9.2)
    add_callout(doc, "NULL is not zero", "NULL wholesale price means no wholesale price was provided. A value of 0.00 means a wholesale price exists and is zero.", fill=PALE_GOLD, accent="7A5A00")

    add_heading(doc, "10. Business and Data Rules", 1)
    add_simple_list(doc, [
        "Every product belongs to exactly one store.",
        "One store can own many products.",
        "Product names are displayed in uppercase for consistent searching and sorting.",
        "Retail, cost, wholesale, and bottle-deposit values cannot be negative.",
        "A barcode may be empty, but a non-empty barcode should not be duplicated within the same store.",
        "A category name should not be duplicated within the same store.",
        "The same person cannot have two membership rows for the same store.",
        "A store supports no more than five active devices in Version 1.",
        "Every saved price change records the new price, updater, and update time.",
        "New products do not show an incorrect previous price of 0.00.",
        "Products should normally be archived instead of permanently deleted.",
        "Different stores must never see or change each other's information.",
    ])

    add_heading(doc, "11. Confirmed and Recommended Decisions", 1)
    decisions = [
        ("Use UUID primary keys", "Confirmed", "Already used by stores and products."),
        ("Products belong to stores through store_id", "Confirmed", "Foreign key already exists."),
        ("Wholesale price remains optional", "Confirmed", "Matches the Version 1 product scope."),
        ("Separate price history from current product price", "Recommended", "Needed for Price Pulse and audit history."),
        ("Create a separate categories table", "Recommended", "Avoids spelling differences and supports category management."),
        ("Separate profiles from stores", "Recommended", "Supports multiple family members cleanly."),
        ("Use store_members to connect users and stores", "Recommended", "Provides a foundation for roles and RLS."),
        ("Track device nickname separately", "Recommended", "Supports update attribution and device limits."),
        ("Archive products instead of immediate deletion", "Recommended", "Protects price history and reduces accidental data loss."),
    ]
    add_table(doc, ["Decision", "Status", "Reason"], decisions, [3500, 1600, 4260], size=9)

    add_heading(doc, "12. Unresolved Questions", 1)
    questions = [
        ("Authentication method", "Will Version 1 use mobile OTP, email/password, or another Supabase Auth method?", "Phase 8"),
        ("Store PIN", "Will the interface keep a PIN-like experience, and how will it be protected server-side?", "Phases 8-9"),
        ("Roles", "Can ordinary members add products, or only update prices?", "Before Phase 9"),
        ("Product deletion", "Who may archive or permanently delete a product?", "Before Phase 9"),
        ("Category ownership", "Are the nine default categories fixed, editable, or both?", "Before Phase 7"),
        ("Barcode duplication", "Should duplicate barcodes be warned about or completely blocked?", "Before Phase 7"),
        ("Offline conflicts", "What happens when two devices change the same price while offline?", "Later synchronization design"),
        ("Store deletion", "Should deleting a store be prohibited, archived, or require a protected recovery period?", "Before production release"),
    ]
    add_table(doc, ["Question", "Decision needed", "When to decide"], questions, [2250, 5100, 2010], size=8.8)

    add_heading(doc, "13. Decision Guide", 1)
    add_body(doc, "Use the following process when an unresolved question appears:")
    add_simple_list(doc, [
        "Write the user problem in plain language.",
        "List the smallest Version 1 behavior that solves it.",
        "Check whether the choice affects privacy, security, or data loss.",
        "Identify the entities, fields, and relationships affected.",
        "Choose one behavior and record the reason.",
        "Move future improvements into a later-version list instead of expanding Version 1.",
    ], numbered=True)

    add_heading(doc, "14. Phase 7 Handoff Guide", 1)
    add_callout(doc, "Do not run SQL yet", "Review and approve this model first. Phase 7 turns only approved entities and rules into actual Supabase tables and constraints.")
    add_simple_list(doc, [
        "Confirm the entity list and fields.",
        "Resolve category behavior and role permissions that affect table structure.",
        "Back up or export the current public schema.",
        "Create one reviewed migration or SQL script for each related group of tables.",
        "Build profiles and store_members before authentication policies depend on them.",
        "Build categories, then update products to use category_id.",
        "Build devices and price_updates.",
        "Add indexes and constraints.",
        "Test with sample data before connecting real store information.",
        "Begin RLS only after ownership relationships are correct.",
    ], numbered=True)

    add_heading(doc, "15. Phase 6 Completion Checklist", 1)
    checks = [
        ("Entity list reviewed", "Pending"),
        ("Every Version 1 field belongs to an entity", "Pending"),
        ("Required and optional rules reviewed", "Pending"),
        ("Primary keys approved", "Pending"),
        ("Foreign keys approved", "Pending"),
        ("One-to-many and one-to-one relationships reviewed", "Pending"),
        ("Price history model approved", "Pending"),
        ("Store ownership of all shared data is clear", "Pending"),
        ("Critical data-structure questions resolved", "Pending"),
        ("Document marked approved for Phase 7", "Pending"),
    ]
    add_table(doc, ["Completion item", "Status"], checks, [7300, 2060], header_fill=PALE_GREEN, size=9.5)

    add_heading(doc, "16. Approval Record", 1)
    approval = [
        ("Document version", "Draft 1.0"),
        ("Phase", "Phase 6 - Data Modeling"),
        ("Approved for Phase 7?", "No - review required"),
        ("Reviewed by", ""),
        ("Review date", ""),
        ("Important changes requested", ""),
    ]
    add_table(doc, ["Item", "Your input"], approval, [2600, 6760], header_fill=PALE_BLUE, size=9.7)

    doc.core_properties.title = "SariSync Phase 6 Data Model"
    doc.core_properties.subject = "Version 1 logical data model and database planning guide"
    doc.core_properties.author = "SariSync Project"
    doc.core_properties.keywords = "SariSync, data model, Supabase, Phase 6, ERD"
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
