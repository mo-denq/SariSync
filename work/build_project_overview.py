# FILE: work/build_project_overview.py | FOLDER: work
# PURPOSE: Builds the editable SariSync project-overview DOCX from documented
# goals, scope, architecture, and implementation notes.
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(r"C:\Users\Yasmien\OneDrive - sdca.edu.ph\Documents\Digital Price Book")
OUT = ROOT / "outputs" / "SariSync Project Overview.docx"
OUT.parent.mkdir(parents=True, exist_ok=True)

# standard_business_brief preset, with a restrained customer_pack opening block
NAVY = "17324D"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
ORANGE = "E8793E"
INK = "26323B"
MUTED = "66727D"
PALE_BLUE = "E8EEF5"
PALE_ORANGE = "FFF2E8"
PALE_GREEN = "EAF5EF"
LIGHT_GRAY = "F2F4F7"
WHITE = "FFFFFF"
BORDER = "D7DEE5"
GREEN = "287A56"

doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(1)
section.right_margin = Inches(1)
section.bottom_margin = Inches(1)
section.left_margin = Inches(1)
section.header_distance = Inches(0.492)
section.footer_distance = Inches(0.492)

def set_font(run, size=11, bold=False, color=INK, italic=False, name="Calibri"):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rpr.rFonts.set(qn("w:ascii"), name)
    rpr.rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)
    return run

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(11)
normal.font.color.rgb = RGBColor.from_string(INK)
normal.paragraph_format.space_before = Pt(0)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.10

for style_name, size, color, before, after in [
    ("Title", 29, NAVY, 0, 8),
    ("Subtitle", 13.5, MUTED, 0, 16),
    ("Heading 1", 16, BLUE, 16, 8),
    ("Heading 2", 13, BLUE, 12, 6),
    ("Heading 3", 12, DARK_BLUE, 8, 4),
]:
    s = styles[style_name]
    s.font.name = "Calibri"
    s.font.size = Pt(size)
    s.font.color.rgb = RGBColor.from_string(color)
    s.font.bold = style_name != "Subtitle"
    s.paragraph_format.space_before = Pt(before)
    s.paragraph_format.space_after = Pt(after)
    s.paragraph_format.keep_with_next = True

# Explicit real-list formatting for preset compliance.
for style_name in ["List Bullet", "List Number"]:
    s = styles[style_name]
    s.font.name = "Calibri"
    s.font.size = Pt(11)
    s.paragraph_format.left_indent = Inches(0.5)
    s.paragraph_format.first_line_indent = Inches(-0.25)
    s.paragraph_format.space_after = Pt(8)
    s.paragraph_format.line_spacing = 1.167

def shade(cell, color):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = tcpr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcpr.append(shd)
    shd.set(qn("w:fill"), color)

def cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tcpr = cell._tc.get_or_add_tcPr()
    mar = tcpr.first_child_found_in("w:tcMar")
    if mar is None:
        mar = OxmlElement("w:tcMar")
        tcpr.append(mar)
    for edge, value in [("top", top), ("start", start), ("bottom", bottom), ("end", end)]:
        item = mar.find(qn(f"w:{edge}"))
        if item is None:
            item = OxmlElement(f"w:{edge}")
            mar.append(item)
        item.set(qn("w:w"), str(value))
        item.set(qn("w:type"), "dxa")

def borders(table, color=BORDER, size="5"):
    tblpr = table._tbl.tblPr
    node = tblpr.find(qn("w:tblBorders"))
    if node is None:
        node = OxmlElement("w:tblBorders")
        tblpr.append(node)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        line = OxmlElement(f"w:{edge}")
        line.set(qn("w:val"), "single")
        line.set(qn("w:sz"), size)
        line.set(qn("w:color"), color)
        node.append(line)

def set_table_width(table, widths_inches, indent_dxa=120):
    widths = [int(x * 1440) for x in widths_inches]
    total = sum(widths)
    tblpr = table._tbl.tblPr
    tblw = tblpr.find(qn("w:tblW"))
    if tblw is None:
        tblw = OxmlElement("w:tblW")
        tblpr.append(tblw)
    tblw.set(qn("w:w"), str(total)); tblw.set(qn("w:type"), "dxa")
    ind = tblpr.find(qn("w:tblInd"))
    if ind is None:
        ind = OxmlElement("w:tblInd"); tblpr.append(ind)
    ind.set(qn("w:w"), str(indent_dxa)); ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid): grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol"); col.set(qn("w:w"), str(width)); grid.append(col)
    for row in table.rows:
        for i, cell in enumerate(row.cells):
            cell.width = Inches(widths_inches[i])
            tcw = cell._tc.get_or_add_tcPr().find(qn("w:tcW"))
            if tcw is None:
                tcw = OxmlElement("w:tcW"); cell._tc.get_or_add_tcPr().append(tcw)
            tcw.set(qn("w:w"), str(widths[i])); tcw.set(qn("w:type"), "dxa")

def cell_text(cell, value, bold=False, color=INK, size=9.4, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.05
    set_font(p.add_run(str(value)), size=size, bold=bold, color=color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    cell_margins(cell)

def add_table(headers, rows, widths, font_size=9.3):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for i, header in enumerate(headers):
        cell_text(t.rows[0].cells[i], header, True, WHITE, 9.2)
        shade(t.rows[0].cells[i], NAVY)
    for row_index, row in enumerate(rows):
        cells = t.add_row().cells
        for i, value in enumerate(row):
            cell_text(cells[i], value, False, INK, font_size)
            if row_index % 2 == 1:
                shade(cells[i], "F8FAFB")
    set_table_width(t, widths)
    borders(t)
    header_props = t.rows[0]._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader"); repeat.set(qn("w:val"), "true"); header_props.append(repeat)
    for row in t.rows:
        cant_split = OxmlElement("w:cantSplit")
        row._tr.get_or_add_trPr().append(cant_split)
    spacer = doc.add_paragraph(); spacer.paragraph_format.space_after = Pt(1)
    return t

def para(text="", size=11, bold=False, color=INK, italic=False, align=None, after=6):
    p = doc.add_paragraph()
    if align is not None: p.alignment = align
    p.paragraph_format.space_after = Pt(after)
    set_font(p.add_run(text), size=size, bold=bold, color=color, italic=italic)
    return p

def bullet(text):
    p = doc.add_paragraph(style="List Bullet")
    set_font(p.add_run(text), size=11)
    return p

def heading(text, level=1):
    return doc.add_heading(text, level=level)

def callout(label, text, fill=PALE_BLUE, accent=BLUE):
    # Paragraph-based note box: avoids using a layout table for prose.
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.12)
    p.paragraph_format.right_indent = Inches(0.12)
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(9)
    p.paragraph_format.line_spacing = 1.10
    ppr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd"); shd.set(qn("w:fill"), fill); ppr.append(shd)
    pbdr = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single"); left.set(qn("w:sz"), "18"); left.set(qn("w:color"), accent); left.set(qn("w:space"), "8")
    pbdr.append(left); ppr.append(pbdr)
    set_font(p.add_run(label.upper() + "\n"), size=9.5, bold=True, color=accent)
    set_font(p.add_run(text), size=11, color=INK)

def page_break():
    doc.add_page_break()

# Running furniture
header = section.header.paragraphs[0]
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
set_font(header.add_run("SARISYNC  |  PROJECT OVERVIEW"), size=8.5, bold=True, color=MUTED)
footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_font(footer.add_run("Shared Digital Price Book for Family-Run Sari-Sari Stores"), size=8.5, color=MUTED)

# Customer-pack opening block
para("PROJECT OVERVIEW", 10, True, ORANGE, align=WD_ALIGN_PARAGRAPH.LEFT, after=3)
para("SariSync", 31, True, NAVY, align=WD_ALIGN_PARAGRAPH.LEFT, after=4)
para("A shared digital price book for small and family-run sari-sari stores", 14, False, MUTED, align=WD_ALIGN_PARAGRAPH.LEFT, after=18)

add_table(
    ["Document purpose", "Current project stage"],
    [["Define the problem, users, purpose, goals, scope, and measures of success.",
      "HTML/CSS/JavaScript prototype using an existing in-app page system; Supabase integration is planned and partially configured."]],
    [3.25, 3.25], 9.6
)

callout(
    "Project in one sentence",
    "SariSync helps family members quickly find the latest selling price of a product, keep prices synchronized across devices, and see who changed a price and when.",
    PALE_ORANGE, ORANGE
)

heading("1. Executive Summary", 1)
para(
    "SariSync is a focused digital price book designed for small and family-run sari-sari stores. It was conceived after observing that fluctuating grocery prices and reliance on memory could cause accidental underpricing and lost profit. Existing store applications, including full point-of-sale tools such as Peddlr, offered more functions than the family needed and introduced too much complexity for users who wanted a fast way to check or update a price."
)
para(
    "The proposed solution is intentionally narrower than a POS system. It centers on rapid product lookup, simple price maintenance, CSV-based bulk setup, shared family access, real-time updates through Supabase, and an audit trail showing which family member made a change. The current implementation is a browser-based HTML/CSS/JavaScript prototype with centralized in-app page navigation."
)

heading("2. Core Problem", 1)
callout(
    "Problem statement",
    "The family sari-sari store was losing money because grocery prices changed frequently, while family members relied on memory or outdated price lists. Existing applications were too complex for the intended users, making price maintenance slow and discouraging consistent use.",
    PALE_ORANGE, ORANGE
)

heading("2.1 What is happening", 2)
bullet("Supplier and grocery prices fluctuate, so an item’s correct selling price may change between purchases.")
bullet("Different family members take turns operating the store and may remember different prices.")
bullet("An outdated or incorrectly remembered price can lead to accidental underpricing and reduced profit.")
bullet("Large product lists are difficult to type manually and inconvenient to maintain on several devices.")

heading("2.2 Why current alternatives do not fit", 2)
bullet("Full POS and ledger applications require more data entry and training than the family’s immediate price-checking need.")
bullet("Complex menus and workflows create technology fatigue for parents or occasional store operators.")
bullet("A feature-rich system can slow down the most common counter task: answering, “How much is this item now?”")

heading("2.3 Consequences if the problem remains", 2)
add_table(
    ["Impact area", "Likely consequence"],
    [
        ["Profit", "Items may be sold below the intended price after supplier costs increase."],
        ["Consistency", "Different family members may quote different prices for the same product."],
        ["Trust", "Unclear or unexplained changes can create confusion among store operators."],
        ["Time", "Users repeatedly ask each other for prices or search through paper and spreadsheet lists."],
        ["Data quality", "Separate lists and memories become outdated and difficult to reconcile."],
    ],
    [1.35, 5.15], 9.5
)

page_break()
heading("3. Intended Users", 1)
para("SariSync is designed primarily for small sari-sari stores where family members share responsibility for daily operations. It favors clarity, speed, and low-friction use over enterprise-level complexity.")

add_table(
    ["User group", "Typical needs", "Design implication"],
    [
        ["Store owner or parent", "Maintain selling prices, import the master list, review changes, and protect profit.", "Simple administration, readable prices, and clear confirmation before important actions."],
        ["Family member at the counter", "Find the current price immediately and understand whether it recently changed.", "Prominent search and barcode scan, minimal steps, and large price text."],
        ["Trusted administrator", "Add products, manage categories, update prices, and correct records.", "Efficient forms, validation, permissions, and reliable history."],
        ["Occasional helper", "Look up prices without learning a full POS workflow.", "Familiar navigation, limited complexity, and safe access appropriate to the role."],
    ],
    [1.45, 2.55, 2.5], 8.9
)

heading("4. Purpose and Value Proposition", 1)
heading("4.1 Purpose", 2)
para(
    "The purpose of SariSync is to provide one shared and current source of truth for product prices. It should allow a family member to find an item’s latest price in seconds, update that price with minimal effort, and make the change visible to the rest of the family."
)

heading("4.2 Value proposition", 2)
callout(
    "Focused by design",
    "SariSync does one job well: it keeps a family store’s product prices easy to find, easy to update, and transparent across the people who operate the store.",
    PALE_GREEN, GREEN
)

heading("4.3 Core product principles", 2)
bullet("Price first: the latest selling price is the most important information on product screens.")
bullet("Fast at the counter: search and barcode scanning should require minimal steps.")
bullet("Simple for the family: workflows should remain understandable to users with different levels of technical confidence.")
bullet("Shared and transparent: changes should synchronize and identify the person and time responsible.")
bullet("Reliable under real conditions: the interface should work on budget phones, different screen sizes, and imperfect connectivity.")

heading("5. Project Goals", 1)
add_table(
    ["Goal", "Description", "Expected user benefit"],
    [
        ["Create a digital price book", "Replace fragmented memory and outdated lists with a shared catalogue of current product prices.", "One dependable place to check prices."],
        ["Enable fast lookup", "Allow products to be found by name, category, or barcode with a clear current price.", "Quicker service and fewer pricing mistakes."],
        ["Keep updates synchronized", "Use Supabase so approved changes can reach family members’ devices in near real time.", "Everyone works from the same latest information."],
        ["Support transparency", "Record the family member, time, old price, and new price for each change.", "Clear accountability without a burdensome workflow."],
        ["Reduce setup effort", "Allow the store’s product list to be imported from CSV with validation and useful feedback.", "Faster initial setup and bulk maintenance."],
        ["Remain simple and readable", "Use a focused interface, responsive layout, high-contrast text, and large prices.", "Comfortable use by parents and occasional helpers."],
    ],
    [1.35, 3.05, 2.1], 8.7
)

page_break()
heading("6. Project Scope", 1)
para("The scope is deliberately limited to the shared price-book problem. This boundary protects the project from becoming another complex POS application.")

heading("6.1 Included in Version 1", 2)
add_table(
    ["Capability", "Included content"],
    [
        ["Product catalogue", "Product name, category, barcode, retail price, wholesale price, optional cost information, unit, notes, and active status."],
        ["Fast price lookup", "Search by product name and barcode; open a readable product details view."],
        ["Price maintenance", "Add products and update supported prices with validation and clear save feedback."],
        ["Price history", "Record previous and new values, price type, family member, and timestamp."],
        ["Family profiles", "Identify members by display name or nickname, role, and optional profile photo."],
        ["Categories", "Organize and browse products through a simple category structure."],
        ["CSV import", "Bulk import a product master list with preview, validation, duplicate handling, and error reporting."],
        ["CSV export", "Export store data in a useful backup or portable format."],
        ["Real-time synchronization", "Use Supabase subscriptions so devices receive approved product and price changes."],
        ["Responsive web experience", "Support small phones, larger phones, and tablets through the existing web page system."],
        ["Offline lookup target", "Cache the most recent product data so basic lookup can remain available when connectivity drops."],
    ],
    [1.65, 4.85], 9.1
)

heading("6.2 Explicitly outside Version 1", 2)
bullet("Full sales recording or point-of-sale checkout")
bullet("Inventory quantity tracking and automatic stock deduction")
bullet("Expense, debt, accounting, or ledger management")
bullet("Employee payroll and scheduling")
bullet("Multi-branch store management")
bullet("AI-based price recommendations")
bullet("Advanced business intelligence and financial reporting")

heading("6.3 Possible future scope", 2)
para("After the single-store price-book workflow is reliable, later releases may consider multi-store support, more advanced offline editing, inventory integration, supplier price comparison, richer analytics, or installation as a Progressive Web App. These ideas should not delay Version 1.")

heading("7. High-Level Solution", 1)
add_table(
    ["Solution area", "Planned approach"],
    [
        ["Interface", "Responsive HTML/CSS views with JavaScript-controlled in-app navigation."],
        ["Shared app shell", "One centralized region for Home, Browse, Updates, Settings, Profile, and product workflows to prevent duplicated navigation state."],
        ["Data and authentication", "Supabase for database access, user authentication, store membership, storage, and real-time subscriptions."],
        ["Bulk setup", "CSV import to populate or update the product catalogue efficiently."],
        ["Local resilience", "A local cache for recent product data, with a staged path toward safer offline support."],
        ["Auditability", "A price-history record created whenever a tracked price changes."],
    ],
    [1.55, 4.95], 9.2
)

page_break()
heading("8. Key User Experiences", 1)
heading("8.1 Check a price", 2)
para("A family member opens SariSync, searches for a product or scans its barcode, and sees the current retail and wholesale price clearly. The interaction should be fast enough for a customer waiting at the counter.")

heading("8.2 Update a price", 2)
para("An authorized member opens the product, enters the new price, reviews the change, and saves. SariSync updates the product, records the old and new values with the member’s identity and time, and synchronizes the change to other active devices.")

heading("8.3 Set up the catalogue", 2)
para("The owner imports an existing spreadsheet as CSV. Before saving, SariSync previews the data, validates required fields and prices, identifies duplicate barcodes or products, and provides a clear result summary.")

heading("8.4 Review changes", 2)
para("A family member opens Updates to review the latest price changes, including the product, previous price, current price, updater, and timestamp. This makes unexpected changes easier to understand and verify.")

heading("9. Success Criteria", 1)
para("Success should be measured by whether SariSync improves real store work, not merely by whether screens have been completed. Initial targets below should be validated and adjusted during a family pilot.")

add_table(
    ["Success area", "Initial criterion", "How to verify"],
    [
        ["Lookup speed", "A typical user can find a known product and read its current price within 5 seconds.", "Timed usability test with representative products and family members."],
        ["Price accuracy", "No unresolved mismatch exists between the approved master list and active product prices after synchronization.", "Sample audit before and after pilot shifts."],
        ["Update transparency", "Every tracked price change records the product, price type, old value, new value, updater, and time.", "Database and Updates-page verification."],
        ["Real-time consistency", "A saved price change appears on another connected device within 5 seconds under normal connectivity.", "Two-device synchronization test."],
        ["CSV usability", "A valid store CSV imports successfully, while invalid rows are identified without silently corrupting data.", "Controlled valid/invalid import tests."],
        ["Ease of use", "Family users complete lookup and update tasks without developer assistance after a short introduction.", "Observed pilot sessions and feedback."],
        ["Responsive quality", "Core tasks work without clipped controls or horizontal scrolling on supported phones and tablets.", "Device and responsive-browser test matrix."],
        ["Reliability", "Users receive clear loading, success, offline, validation, and failure messages instead of raw technical errors.", "Scenario-based quality review."],
        ["Scope discipline", "Version 1 is usable as a price book without adding POS, accounting, or inventory complexity.", "Requirements and release review."],
    ],
    [1.25, 3.2, 2.05], 8.5
)

heading("10. Assumptions, Constraints, and Risks", 1)
add_table(
    ["Item", "Project implication or response"],
    [
        ["Intermittent internet", "Prioritize cached lookup and communicate sync status honestly; introduce offline editing only after conflict handling is tested."],
        ["Shared store duties", "Keep common workflows short and preserve an audit trail tied to a reliable member identity."],
        ["Budget devices", "Limit heavy animations and unnecessary images; optimize DOM updates and keep touch targets large."],
        ["Barcode hardware differences", "Provide typed search and manual barcode entry when scanning is unavailable or permission is denied."],
        ["Incorrect CSV data", "Require preview and validation; report rejected rows and define duplicate behavior."],
        ["Simultaneous updates", "Use timestamps or version checks and define a conflict rule before calling the system fully synchronized."],
        ["Technology confidence", "Use plain language, predictable navigation, and strong visual hierarchy rather than feature-dense screens."],
    ],
    [1.55, 4.95], 9.1
)

page_break()
heading("11. Current Project Status", 1)
para("The project currently has a working web-prototype foundation. This overview separates existing implementation from planned production capabilities.")
add_table(
    ["Area", "Current status"],
    [
        ["Frontend", "HTML, CSS, and JavaScript prototype with a shared application shell and responsive styling."],
        ["In-app pages", "Home, Browse, Updates, Settings, Profile, and product workflows are designed to render inside a centralized secondary-page region while lightweight."],
        ["Feature folders", "Home, authentication, and category modules exist; shared components, models, and constants folders are prepared for growth."],
        ["Data", "Prototype product/category data is currently separated through service files; final Supabase data contracts still require completion."],
        ["Supabase", "The JavaScript client dependency and configuration service are present; complete schema, authentication, policies, and real-time behavior remain development work."],
        ["Development emphasis", "The UI and interaction concept are further developed than authentication, database security, offline behavior, and production testing."],
    ],
    [1.55, 4.95], 9.2
)

heading("12. Project Boundaries and Decision Rule", 1)
callout(
    "Decision rule",
    "A proposed feature belongs in Version 1 only if it directly improves price lookup, price maintenance, synchronization, transparency, or safe setup of the shared price book. Otherwise, record it for a future release.",
    PALE_GREEN, GREEN
)

heading("13. Definition of Project Success", 1)
para(
    "SariSync will be successful when family members can rely on it during real store operations: they can find the latest price quickly, make authorized updates without confusion, see who changed a price, and trust that other devices will receive the same information. The app should reduce pricing mistakes without forcing the family to operate a full POS or accounting system."
)

heading("14. Recommended Supporting Documents", 1)
add_table(
    ["Document", "Purpose"],
    [
        ["SariSync Version 1 Requirements", "Translate this overview into testable functional and non-functional requirements."],
        ["UI/UX Specification", "Document pages, navigation, responsive rules, components, and interface states."],
        ["Technical and Database Design", "Define the application architecture, data model, schema, authentication, permissions, and synchronization."],
        ["Testing Plan", "Specify unit, DOM, integration, security, device, offline, CSV, and family-pilot testing."],
        ["Development Roadmap", "Track phases, dependencies, milestones, completion criteria, and current status."],
    ],
    [2.15, 4.35], 9.4
)

doc.core_properties.title = "SariSync Project Overview"
doc.core_properties.subject = "Problem, users, purpose, goals, scope, and success criteria"
doc.core_properties.author = "SariSync Project"
doc.core_properties.keywords = "SariSync, sari-sari store, digital price book, project overview"
doc.save(OUT)
print(OUT)
