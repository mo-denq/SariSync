# FILE: work/qa_phase5_doc.py
# PURPOSE: Performs structural checks on the generated Phase 5 DOCX when a
# visual LibreOffice render is unavailable in the workspace.

from pathlib import Path
from zipfile import ZipFile

from docx import Document
from lxml import etree


DOCX = Path(__file__).resolve().parents[1] / "SariSync_Phase_5_Web_Architecture.docx"
NS = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}


def main():
    document = Document(DOCX)
    all_text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    all_text += "\n" + "\n".join(
        " | ".join(cell.text for cell in row.cells)
        for table in document.tables
        for row in table.rows
    )

    placeholder_terms = ["[[TODO]]", "[[PLACEHOLDER]]", "[[REF:", "Lorem ipsum"]
    placeholder_hits = [term for term in placeholder_terms if term.lower() in all_text.lower()]

    with ZipFile(DOCX) as archive:
        document_xml = etree.fromstring(archive.read("word/document.xml"))
        page_breaks = len(document_xml.xpath(".//w:br[@w:type='page']", namespaces=NS))
        table_header_rows = len(document_xml.xpath(".//w:tblHeader", namespaces=NS))
        table_widths = document_xml.xpath(".//w:tblPr/w:tblW/@w:w", namespaces=NS)
        table_indents = document_xml.xpath(".//w:tblPr/w:tblInd/@w:w", namespaces=NS)
        image_descriptions = document_xml.xpath(".//wp:docPr/@descr", namespaces={
            **NS,
            "wp": "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing",
        })

    bad_table_widths = [width for width in table_widths if width != "9360"]
    bad_table_indents = [indent for indent in table_indents if indent != "120"]
    headings = [
        (paragraph.style.name, paragraph.text)
        for paragraph in document.paragraphs
        if paragraph.style.name.startswith("Heading")
    ]

    print(f"file={DOCX}")
    print(f"bytes={DOCX.stat().st_size}")
    print(f"paragraphs={len(document.paragraphs)}")
    print(f"tables={len(document.tables)}")
    print(f"sections={len(document.sections)}")
    print(f"inline_shapes={len(document.inline_shapes)}")
    print(f"headings={len(headings)}")
    print(f"manual_page_breaks={page_breaks}")
    print(f"table_header_rows={table_header_rows}")
    print(f"bad_table_widths={bad_table_widths}")
    print(f"bad_table_indents={bad_table_indents}")
    print(f"image_descriptions={image_descriptions}")
    print(f"placeholder_hits={placeholder_hits}")
    print("heading_list:")
    for style, text in headings:
        print(f"  {style}: {text}")

    if placeholder_hits or bad_table_widths or bad_table_indents:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
