# FILE: work/build_phase5_web_architecture.py
# PURPOSE: Generates the SariSync Phase 5 Web Architecture document.
# The document explains the current HTML/CSS/JavaScript prototype, the approved
# target structure, responsibility boundaries, staged refactoring, and the
# checks required before Phase 7 database implementation begins.

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "SariSync_Phase_5_Web_Architecture.docx"
ASSET_DIR = ROOT / "work" / "phase5_assets"
ARCH_DIAGRAM = ASSET_DIR / "phase5_architecture_layers.png"

NAVY = "173C3B"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
ORANGE = "F25A24"
PALE_ORANGE = "FFF0E8"
PALE_BLUE = "E8EEF5"
PALE_GREEN = "E8F5EC"
PALE_GOLD = "FFF6D8"
LIGHT = "F7F8F6"
GRID = "D7DDD9"
MUTED = "667673"
WHITE = "FFFFFF"
BLACK = "172C2B"
RED = "A13D32"

PAGE_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120


def set_run(run, size=11, bold=False, color=BLACK, font="Calibri", italic=False):
    run.font.name = font
    r_pr = run._element.get_or_add_rPr()
    r_pr.rFonts.set(qn("w:ascii"), font)
    r_pr.rFonts.set(qn("w:hAnsi"), font)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_table_geometry(table, widths):
    if sum(widths) != PAGE_WIDTH_DXA:
        raise ValueError(f"Table widths must sum to {PAGE_WIDTH_DXA}: {widths}")
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl = table._tbl
    tbl_pr = tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(PAGE_WIDTH_DXA))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT_DXA))
    tbl_ind.set(qn("w:type"), "dxa")

    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths[index]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)


def set_cell_text(cell, text, bold=False, color=BLACK, size=9.2, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = align
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.12
    run = paragraph.add_run(str(text))
    set_run(run, size=size, bold=bold, color=color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(doc, headers, rows, widths, header_fill=PALE_BLUE, size=9.2, center_columns=None):
    center_columns = set(center_columns or [])
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for index, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[index], header, bold=True, color=NAVY, size=9.2)
        set_cell_shading(table.rows[0].cells[index], header_fill)
    set_repeat_table_header(table.rows[0])

    for row_data in rows:
        row = table.add_row()
        for index, value in enumerate(row_data):
            align = WD_ALIGN_PARAGRAPH.CENTER if index in center_columns else WD_ALIGN_PARAGRAPH.LEFT
            set_cell_text(row.cells[index], value, size=size, align=align)
    set_table_geometry(table, widths)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)
    return table


def add_heading(doc, text, level=1):
    paragraph = doc.add_paragraph(text, style=f"Heading {level}")
    paragraph.paragraph_format.keep_with_next = True
    return paragraph


def add_body(doc, text, bold_lead=None):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.25
    if bold_lead and text.startswith(bold_lead):
        lead = paragraph.add_run(bold_lead)
        set_run(lead, bold=True)
        remainder = paragraph.add_run(text[len(bold_lead):])
        set_run(remainder)
    else:
        run = paragraph.add_run(text)
        set_run(run)
    return paragraph


def add_callout(doc, label, text, fill=PALE_ORANGE, accent=ORANGE):
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    set_table_geometry(table, [PAGE_WIDTH_DXA])
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    cell.text = ""
    title = cell.paragraphs[0]
    title.paragraph_format.space_after = Pt(2)
    title_run = title.add_run(label.upper())
    set_run(title_run, size=9, bold=True, color=accent)
    body = cell.add_paragraph()
    body.paragraph_format.space_after = Pt(0)
    body.paragraph_format.line_spacing = 1.2
    body_run = body.add_run(text)
    set_run(body_run, size=10.5)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(1)


def add_list(doc, items, numbered=False, level=0):
    style = "List Number" if numbered else "List Bullet"
    if level == 1:
        style += " 2"
    for item in items:
        paragraph = doc.add_paragraph(style=style)
        paragraph.paragraph_format.space_after = Pt(4)
        paragraph.paragraph_format.line_spacing = 1.25
        run = paragraph.add_run(item)
        set_run(run)


def add_code_block(doc, text, label=None):
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    set_table_geometry(table, [PAGE_WIDTH_DXA])
    cell = table.cell(0, 0)
    set_cell_shading(cell, "F2F4F5")
    cell.text = ""
    if label:
        p_label = cell.paragraphs[0]
        p_label.paragraph_format.space_after = Pt(3)
        r_label = p_label.add_run(label.upper())
        set_run(r_label, size=8.5, bold=True, color=MUTED)
        p = cell.add_paragraph()
    else:
        p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.05
    r = p.add_run(text)
    set_run(r, size=8.4, font="Consolas", color=BLACK)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    set_run(run, size=9, color=MUTED)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    paragraph._p.append(field)


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(BLACK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, DARK_BLUE, 10, 5),
    ):
        style = doc.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for list_name in ("List Bullet", "List Number", "List Bullet 2", "List Number 2"):
        style = doc.styles[list_name]
        style.font.name = "Calibri"
        style.font.size = Pt(11)
        if list_name.endswith(" 2"):
            style.paragraph_format.left_indent = Inches(0.625)
            style.paragraph_format.first_line_indent = Inches(-0.188)
        else:
            style.paragraph_format.left_indent = Inches(0.375)
            style.paragraph_format.first_line_indent = Inches(-0.188)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.25

    header = section.header
    header_paragraph = header.paragraphs[0]
    header_paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    header_run = header_paragraph.add_run("SariSync | Phase 5 Web Architecture")
    set_run(header_run, size=9, bold=True, color=MUTED)

    footer = section.footer
    add_page_number(footer.paragraphs[0])


def create_architecture_diagram():
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    width, height = 1700, 1120
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    try:
        title_font = ImageFont.truetype("arialbd.ttf", 43)
        box_font = ImageFont.truetype("arialbd.ttf", 29)
        sub_font = ImageFont.truetype("arial.ttf", 21)
        arrow_font = ImageFont.truetype("arialbd.ttf", 26)
    except OSError:
        title_font = box_font = sub_font = arrow_font = ImageFont.load_default()

    draw.text((75, 35), "SariSync Target Web Architecture", fill=f"#{NAVY}", font=title_font)
    layers = [
        ("Application shell and page modules", "index.html, navigation.js, pages/*, components/*", 130, PALE_ORANGE),
        ("Central state", "session, active store, navigation stack, cached records, UI status", 305, PALE_BLUE),
        ("Models and services", "validation, barcode, CSV, sync, profile, backup", 480, LIGHT),
        ("Repositories", "stable Product, Category, Profile, and Update data interfaces", 655, PALE_GREEN),
        ("Storage and production data", "local cache / pending queue  |  Supabase through repositories", 830, PALE_GOLD),
    ]

    for index, (title, subtitle, y, fill) in enumerate(layers):
        x1, x2 = 150, 1550
        y1, y2 = y, y + 120
        draw.rounded_rectangle((x1, y1, x2, y2), radius=25, fill=f"#{fill}", outline=f"#{GRID}", width=3)
        draw.text((205, y1 + 23), title, fill=f"#{NAVY}", font=box_font)
        draw.text((205, y1 + 70), subtitle, fill=f"#{MUTED}", font=sub_font)
        if index < len(layers) - 1:
            center_x = width // 2
            draw.line((center_x, y2 + 4, center_x, y2 + 48), fill=f"#{ORANGE}", width=5)
            draw.polygon(
                [(center_x, y2 + 55), (center_x - 13, y2 + 36), (center_x + 13, y2 + 36)],
                fill=f"#{ORANGE}",
            )

    draw.text(
        (150, 1030),
        "Rule: pages request actions; repositories own data access; storage details do not leak into UI modules.",
        fill=f"#{MUTED}",
        font=sub_font,
    )
    image.save(ARCH_DIAGRAM, dpi=(180, 180))


def add_cover(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(95)
    p.paragraph_format.space_after = Pt(13)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("PHASE 5")
    set_run(r, size=11, bold=True, color=ORANGE)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run("SariSync Version 1\nWeb Architecture")
    set_run(r, size=29, bold=True, color=NAVY)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(35)
    r = p.add_run("A maintainable structure for the working HTML, CSS, and JavaScript prototype")
    set_run(r, size=13, color=MUTED, italic=True)

    add_callout(
        doc,
        "Architecture decision",
        "Keep one shared application shell and bottom navigation. Split page rendering, state, reusable components, and data access into focused modules without changing the current UI or user flows.",
        fill=PALE_ORANGE,
    )

    metadata = [
        ("Document version", "Draft 1.0"),
        ("Prepared", "July 29, 2026"),
        ("Current platform", "Browser prototype using HTML, CSS, and JavaScript"),
        ("Related plans", "Phase 3 screen inventory, Version 1 requirements, Phase 6 data model"),
    ]
    add_table(doc, ["Item", "Value"], metadata, [2700, 6660], header_fill=PALE_BLUE, size=9.6)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(28)
    r = p.add_run("Prepared for the SariSync shared digital price book project")
    set_run(r, size=10, color=MUTED)
    doc.add_page_break()


def add_contents(doc):
    add_heading(doc, "Document Map", 1)
    rows = [
        ("1-3", "Purpose, status, current prototype"),
        ("4-6", "Architecture principles, layers, and target folder tree"),
        ("7-10", "Files, shell, navigation, state, pages, and components"),
        ("11-13", "CSS, repositories, services, storage, and security boundaries"),
        ("14-17", "Refactoring sequence, tests, completion checklist, and approval"),
        ("Appendix", "Glossary and evidence reviewed"),
    ]
    add_table(doc, ["Sections", "What they explain"], rows, [1800, 7560], header_fill=PALE_GREEN, size=9.7)
    add_callout(
        doc,
        "How to use this document",
        "Approve the architecture before large file moves. Phase 5 describes responsibilities and safe refactoring order. It does not authorize database migrations or a redesign of the working interface.",
        fill=PALE_GOLD,
        accent="7A5A00",
    )


def build_document():
    create_architecture_diagram()
    doc = Document()
    configure_document(doc)
    add_cover(doc)
    add_contents(doc)

    add_heading(doc, "1. Phase 5 Purpose and Scope", 1)
    add_body(
        doc,
        "Phase 5 defines how the browser prototype is organized so pages, navigation, state, reusable interface pieces, and data operations can change independently. The goal is clarity and maintainability, not a framework migration.",
    )
    add_heading(doc, "1.1 Goals", 2)
    add_list(
        doc,
        [
            "Keep the existing SariSync branding, screens, navigation flow, and responsive behavior.",
            "Give every file one clear primary responsibility.",
            "Keep one application shell responsible for bottom navigation and shared page space.",
            "Move page-specific rendering out of the Home controller.",
            "Provide stable repository interfaces so prototype data can later be replaced by Supabase.",
            "Separate local caching and pending offline edits from page code.",
            "Make the code understandable to a first-time app developer through clear names and file headers.",
        ],
    )
    add_heading(doc, "1.2 Non-goals", 2)
    add_list(
        doc,
        [
            "Do not redesign the approved interface.",
            "Do not replace the browser prototype with Flutter, React Native, or another framework in Phase 5.",
            "Do not create or migrate Supabase tables; that belongs to Phase 7.",
            "Do not finalize authentication policies or Row Level Security; those belong to later phases.",
            "Do not promise production-grade offline synchronization until its conflict behavior is designed.",
        ],
    )
    add_callout(
        doc,
        "Phase boundary",
        "Phase 5 produces an approved architecture and staged refactoring plan. Phase 6 defines records and relationships. Phase 7 implements the approved database design.",
        fill=PALE_BLUE,
        accent=BLUE,
    )

    add_heading(doc, "2. Current Phase 5 Status", 1)
    status_rows = [
        ("Shared application shell", "Exists", "index.html owns the Home shell, secondary view region, bottom navigation, floating add action, sheet, and toast."),
        ("Feature folders", "Partly exists", "Auth and category pages are separated; most other views remain in home.js."),
        ("Shared services", "Partly exists", "Barcode, responsive helpers, formatting, prototype data, and Supabase client setup are separate."),
        ("Reusable components", "Reserved only", "components/ contains guidance but no component modules yet."),
        ("Models", "Reserved only", "models/ contains guidance; product and category shapes still live in services/data.js."),
        ("Central state", "Not implemented", "State is spread across variables, global window objects, and localStorage."),
        ("Repositories", "Not implemented", "Page code reads and mutates prototype arrays directly."),
        ("CSS layers", "Not implemented", "Most shared and page-specific styling remains in one styles.css file."),
        ("Automated tests", "Not implemented", "No focused unit or architecture tests are present."),
    ]
    add_table(doc, ["Area", "Status", "Observed meaning"], status_rows, [2500, 1650, 5210], header_fill=PALE_GREEN, size=8.9, center_columns={1})
    add_callout(
        doc,
        "Current position",
        "The prototype already demonstrates the user experience, but its internal organization is still transitional. The next work is an incremental separation of responsibilities, not a destructive rewrite.",
    )

    add_heading(doc, "3. Current Prototype Assessment", 1)
    current_rows = [
        ("index.html", "Shared app shell and Home markup", "Correct direction; retain as the single shell."),
        ("pages/home/home.js", "Home, Browse, Updates, Settings, Profile, product details, add/edit, navigation, CSV, and mutations", "Main concentration risk. Split by responsibility."),
        ("styles.css", "All design tokens, base styles, layouts, components, pages, and responsive rules", "Works, but difficult to trace and safely change."),
        ("pages/auth/*", "Login and store-registration flow", "Already separate, but the old PIN flow conflicts with the latest email/password decision."),
        ("pages/categories/*", "Standalone category browser", "Useful reusable page; data transfer currently depends on localStorage."),
        ("services/data.js", "Sample products and categories", "Clearly label as prototype data; later move behind repositories."),
        ("services/barcode.js", "Camera scanning behavior", "Good service boundary; keep page-neutral."),
        ("services/utils.js", "Title-case and peso formatting", "Good shared helper boundary."),
        ("services/responsive.js", "Viewport-derived CSS tokens", "Keep as a focused responsive helper."),
        ("services/supabase.js", "Browser Supabase client", "Keep client initialization separate; do not place secret/service-role credentials in frontend code."),
        ("settings.html", "Compatibility redirect to the shared Settings view", "Retain temporarily, then remove only after old links are no longer needed."),
        ("components/, models/, constants/", "Reserved folders with README guidance", "Populate gradually during refactoring."),
    ]
    add_table(doc, ["Current file or area", "Present responsibility", "Phase 5 conclusion"], current_rows, [2300, 3450, 3610], size=8.4)

    add_heading(doc, "3.1 Main Architectural Risks", 2)
    add_list(
        doc,
        [
            "Changing Browse, Settings, or product forms can unintentionally affect Home because they share one controller.",
            "Page code directly mutates shared product arrays, making later Supabase and offline behavior difficult to introduce safely.",
            "Global window objects hide dependencies and make tests harder to isolate.",
            "localStorage is used for both session-like information and category-page transfer without one storage contract.",
            "The current auth prototype mixes real Supabase calls with prototype session behavior and an outdated PIN path.",
            "One large stylesheet makes it difficult to distinguish a reusable component rule from a page-only exception.",
        ],
    )

    doc.add_page_break()
    add_heading(doc, "4. Architecture Principles", 1)
    principles = [
        ("One shell", "index.html owns shared navigation, shared overlays, and the active page container."),
        ("One-way dependencies", "Pages may use components, state, repositories, and services. Repositories and services must not depend on pages."),
        ("State has one home", "Session, active store, active view, navigation stack, records, and UI status live in a central store."),
        ("Repositories hide storage", "Pages request products and categories through stable interfaces, not directly from arrays, localStorage, or Supabase."),
        ("Components stay presentation-focused", "A product card receives data and callbacks; it does not fetch or save products."),
        ("Services perform operations", "Barcode, CSV, sync, profile photo, and backup logic remain separate from page rendering."),
        ("Models normalize data", "Names, optional prices, barcodes, category values, and timestamps are validated consistently."),
        ("Safe incremental refactoring", "Move one responsibility, verify behavior, then continue."),
    ]
    add_table(doc, ["Principle", "Plain-language rule"], principles, [2400, 6960], size=9.2)
    add_callout(
        doc,
        "Dependency rule",
        "UI code may call downward into state, repositories, and services. Data and service modules must never reach upward to manipulate a specific page.",
        fill=PALE_GOLD,
        accent="7A5A00",
    )

    add_heading(doc, "5. Target Architecture Overview", 1)
    add_body(
        doc,
        "The target architecture keeps the prototype framework-free while introducing clear layers. Each layer exposes a small interface to the layer above it.",
    )
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picture_run = p.add_run()
    picture = picture_run.add_picture(str(ARCH_DIAGRAM), width=Inches(6.45))
    picture._inline.docPr.set(
        "descr",
        "Layered SariSync web architecture showing pages and components, central state, models and services, repositories, and local or Supabase storage.",
    )
    p.paragraph_format.space_after = Pt(5)
    caption = doc.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_after = Pt(8)
    r = caption.add_run("Figure 1. Approved dependency direction for the SariSync browser prototype.")
    set_run(r, size=9, italic=True, color=MUTED)

    layer_rows = [
        ("Shell and pages", "Render views, bind user events, request actions, and show loading/empty/error states."),
        ("Components", "Render reusable interface fragments from supplied data and callbacks."),
        ("Central state", "Holds shared client state and notifies views when relevant data changes."),
        ("Models", "Normalize, validate, serialize, and compare application records."),
        ("Services", "Perform feature operations such as scanning, CSV parsing, sync, profile photos, and backups."),
        ("Repositories", "Provide stable create/read/update/archive interfaces and choose the active data source."),
        ("Storage", "Cache records locally, preserve pending edits, and communicate with Supabase through repositories."),
    ]
    add_table(doc, ["Layer", "Responsibility"], layer_rows, [2100, 7260], size=9.3)

    doc.add_page_break()
    add_heading(doc, "6. Target Folder and File Structure", 1)
    add_body(
        doc,
        "This is the Phase 5 target structure requested for SariSync. It is a destination map, not an instruction to move every file at once.",
    )
    target_tree = """src/
  index.html                 # shared application shell
  css/
    tokens.css               # colors, spacing, radius, type, breakpoints
    base.css                 # reset, typography, accessibility
    layout.css               # shell, page container, navigation
    components.css           # buttons, cards, fields, modals
    pages/                   # view-specific styles only when necessary
  js/
    app.js                   # startup and dependency wiring
    navigation.js            # active tab, secondary views, Back stack
    state/store.js           # central client state
    pages/
      home.js
      browse.js
      updates.js
      settings.js
      profile.js
      product-details.js
    components/
      page-header.js
      product-card.js
      modal-shell.js
      toast.js
    models/                  # normalization and validation helpers
    repositories/            # stable data interface
    services/                # barcode, CSV, sync, profile, backup
    storage/                 # local cache / localStorage adapter
  assets/
  tests/"""
    add_code_block(doc, target_tree, "Approved target structure")

    add_heading(doc, "6.1 SariSync Version 1 Extension Files", 2)
    add_body(
        doc,
        "The requested tree is the minimum. These page modules can be added when their current render functions are extracted from home.js.",
    )
    extensions = [
        ("js/pages/add-product.js", "Add-product page and form orchestration."),
        ("js/pages/update-price.js", "Product finder for the update workflow."),
        ("js/pages/edit-product.js", "Edit product description and price form orchestration."),
        ("js/pages/category-products.js", "Reusable category item view."),
        ("js/pages/import-price-list.js", "CSV selection, preview, validation results, and confirmation."),
        ("js/pages/export-backup.js", "Backup preparation and download status."),
        ("js/pages/auth.js", "Page orchestration for email/password registration and login; implementation remains Phase 8."),
        ("js/components/settings-tile.js", "Reusable Settings option with title, subtitle, and action."),
        ("js/components/empty-state.js", "Reusable no-data presentation and optional action."),
        ("js/components/barcode-button.js", "Shared barcode scan button used by Search and product forms."),
    ]
    add_table(doc, ["Extension file", "Responsibility"], extensions, [3300, 6060], size=8.9)

    add_heading(doc, "6.2 Naming Rules", 2)
    add_list(
        doc,
        [
            "Use lowercase kebab-case file names: product-details.js, not ProductDetails.js.",
            "Name page entry points after the view: browse.js and settings.js.",
            "Name reusable UI after what it renders: product-card.js and page-header.js.",
            "Name services after the operation: barcode-service.js, csv-service.js, and sync-service.js.",
            "Name repositories after the record type: product-repository.js and category-repository.js.",
            "Place a beginner-friendly FILE and PURPOSE comment at the top of every supported source file.",
        ],
    )

    add_heading(doc, "7. File and Folder Responsibilities", 1)
    responsibility_rows = [
        ("src/index.html", "Static shared shell only", "Bottom navigation, page outlet, shared sheet/toast roots, script/style entry points."),
        ("src/js/app.js", "Startup and wiring", "Create dependencies, restore session/cache, start navigation, and render the initial view."),
        ("src/js/navigation.js", "Navigation state", "Activate tabs, open secondary pages, maintain Back stack, and update browser history if used."),
        ("src/js/state/store.js", "Shared client state", "Read, update, and subscribe to session, store, device, records, and UI status."),
        ("src/js/pages/", "Page orchestration", "Render one view, bind its events, call repositories/services, and display states."),
        ("src/js/components/", "Reusable presentation", "Return DOM/markup and expose callbacks without owning page business rules."),
        ("src/js/models/", "Record contracts", "Normalize names, optional fields, money values, timestamps, and IDs."),
        ("src/js/repositories/", "Data access contracts", "Hide whether records currently come from mock arrays, local cache, or Supabase."),
        ("src/js/services/", "Feature operations", "Barcode, CSV, sync, profile image, and backup behavior."),
        ("src/js/storage/", "Browser persistence", "Read/write cache, pending queue, device registration, and safe session preferences."),
        ("src/css/", "Shared design system", "Tokens, accessible defaults, shell layout, components, and exceptional page styles."),
        ("src/assets/", "Static media", "App marks, placeholder images, and non-code assets."),
        ("src/tests/", "Automated checks", "Models, repositories, navigation, services, and critical rendering helpers."),
    ]
    add_table(doc, ["Location", "Primary responsibility", "May contain"], responsibility_rows, [2500, 2500, 4360], size=8.3)

    doc.add_page_break()
    add_heading(doc, "8. Shared Application Shell and Navigation", 1)
    add_heading(doc, "8.1 Shell Ownership", 2)
    add_body(
        doc,
        "index.html remains the single shared shell. Individual page modules render inside one page outlet. Pages must not create another bottom navigation bar, another full application shell, or a nested fixed viewport.",
    )
    shell_owns = [
        "App-level background and maximum-width container.",
        "Home-only branded header host.",
        "Active page outlet for Home, Browse, Updates, Settings, Profile, and workflows.",
        "One opaque bottom navigation bar and floating add action.",
        "Shared toast, confirmation dialog, scanner, and modal/sheet roots.",
        "Safe-area spacing and the final content inset above bottom navigation.",
    ]
    add_list(doc, shell_owns)

    add_heading(doc, "8.2 Navigation Contract", 2)
    navigation_rows = [
        ("Home, Browse, Updates, Settings", "Bottom navigation", "Replace the active primary view; keep the shared shell mounted."),
        ("Profile", "Home avatar or Settings option", "Push Profile onto the Back stack; Back returns to the exact previous view."),
        ("Product details", "Search, Browse, Category, or Price Pulse", "Push details; Back returns to the originating list and preserves its search/filter."),
        ("Add Product", "Floating add action or Quick Actions", "Push page; successful save returns to the intended previous view."),
        ("Update Price", "Quick Action", "Push product finder, then edit page; Back unwinds one step at a time."),
        ("Category Products", "Browse or Home category", "Push category view; Back returns to Browse or Home origin."),
        ("Import / Export", "Settings", "Push administrative view; Back returns to Settings."),
    ]
    add_table(doc, ["Destination", "Entry", "Back behavior"], navigation_rows, [2300, 2800, 4260], size=8.7)
    add_callout(
        doc,
        "Back-stack rule",
        "navigation.js stores view identifiers and serializable parameters, not raw DOM. A page re-renders from state when the user goes Back.",
        fill=PALE_BLUE,
        accent=BLUE,
    )

    add_heading(doc, "8.3 Suggested Navigation Interface", 2)
    add_code_block(
        doc,
        """navigation.start({ initialView: "home" })
navigation.switchTab("browse")
navigation.push("product-details", { productId, from: "browse" })
navigation.back()
navigation.replace("home")""",
        "Illustrative interface",
    )

    add_heading(doc, "9. Central State", 1)
    add_body(
        doc,
        "The central state store replaces scattered page variables and reduces unnecessary rebuilding. A page subscribes only to the data it needs.",
    )
    state_rows = [
        ("session", "Authenticated user ID, signed-in status, and session readiness."),
        ("activeStore", "Store ID, generated store code, display name, and membership role."),
        ("activeDevice", "Device ID, nickname, active status, and last activity."),
        ("navigation", "Active tab, current view, parameters, and Back stack."),
        ("products", "Normalized active and archived product records for the selected store."),
        ("categories", "Store-owned, uppercase category records and display order."),
        ("priceUpdates", "Recent audit records used by Price Pulse and Updates."),
        ("sync", "Online, offline, pending, syncing, conflict, or error state."),
        ("ui", "Search terms, selected filters, loading/error messages, and transient notices."),
    ]
    add_table(doc, ["State branch", "Contains"], state_rows, [2200, 7160], size=9.1)
    add_heading(doc, "9.1 State Rules", 2)
    add_list(
        doc,
        [
            "Pages never mutate state objects directly; they call named actions.",
            "Repository results are normalized before entering state.",
            "Optional numeric values remain null when absent; blank wholesale price is never converted to zero.",
            "Archived products stay available to audit/history features but are excluded from normal active lists.",
            "A product update changes only subscribers interested in that product or price-update collection.",
            "localStorage is accessed through the storage adapter, not directly from page modules.",
        ],
    )

    add_heading(doc, "10. Page and Component Boundaries", 1)
    page_rows = [
        ("home.js", "Greeting, primary search, Price Pulse preview, Quick Actions, and category preview.", "Navigation, repositories, shared components."),
        ("browse.js", "Full category/product browsing, filters, search, and edit entry.", "Product/category repositories and product cards."),
        ("updates.js", "Complete price history and audit trail.", "Price-update repository and price cards."),
        ("settings.js", "Account, store management, import/export, sync, and About entries.", "Settings tiles and navigation."),
        ("profile.js", "Profile display/editing and photo action.", "Profile repository/service and profile components."),
        ("product-details.js", "Current retail/wholesale values, metadata, timestamps, and edit action.", "Product repository and price display components."),
        ("add-product.js", "New-product form and barcode action.", "Product model, repository, barcode service."),
        ("update-price.js", "Search/category selection before editing.", "Product repository and navigation."),
        ("edit-product.js", "Price and description edits for one product.", "Product model, repository, price-update workflow."),
    ]
    add_table(doc, ["Page module", "Owns", "Uses"], page_rows, [2100, 4500, 2760], size=8.2)

    component_rows = [
        ("page-header.js", "Title, optional Back button, optional action.", "No branding outside Home."),
        ("product-card.js", "Product name, category/variant, current price, and select/edit action.", "Receives product and callbacks."),
        ("price-card.js", "Old/new price, direction, updater, device, and time.", "Receives one price-update view model."),
        ("modal-shell.js", "Shared rounded container, header, scroll body, and fixed footer.", "Use only when a modal is genuinely required."),
        ("toast.js", "Temporary success, warning, and error messages.", "One shared live region."),
        ("settings-tile.js", "Settings label, subtitle, and navigation action.", "Presentation only."),
        ("barcode-button.js", "Consistent scan icon, label, and accessible action.", "Calls supplied scan callback."),
        ("empty-state.js", "No-data message and optional recovery action.", "Receives text and callback."),
    ]
    add_heading(doc, "10.1 Reusable Component Contract", 2)
    add_table(doc, ["Component", "Renders", "Boundary"], component_rows, [2200, 4000, 3160], size=8.5)
    add_callout(
        doc,
        "Component rule",
        "Components receive data and callbacks. They do not read localStorage, query Supabase, change active navigation, or mutate the product collection.",
        fill=PALE_GOLD,
        accent="7A5A00",
    )

    doc.add_page_break()
    add_heading(doc, "11. CSS and Responsive Architecture", 1)
    css_rows = [
        ("tokens.css", "Color, spacing, radius, shadows, type scale, content width, safe-area values, and responsive breakpoints."),
        ("base.css", "Box sizing, body defaults, typography, focus visibility, reduced motion, screen-reader helpers, and form inheritance."),
        ("layout.css", "Application shell, page outlet, shared page wrapper, bottom navigation, safe-area handling, grids, and scroll boundaries."),
        ("components.css", "Buttons, fields, cards, chips, tiles, toasts, scanners, dialogs, and shared states."),
        ("pages/*.css", "Only rules unique to one page that cannot be expressed through shared layout or components."),
    ]
    add_table(doc, ["File", "Responsibility"], css_rows, [2200, 7160], size=9.2)

    add_heading(doc, "11.1 Elastic Storefront Rules", 2)
    add_list(
        doc,
        [
            "Use percentage widths and Flexbox/Grid for structural layout; avoid hardcoded card widths.",
            "Center the page wrapper with a shared maximum width so tablets do not overstretch content.",
            "Use clamped numeric CSS custom properties for typography; prices retain a safe minimum and maximum.",
            "Use one shared horizontal-padding token across Home, Browse, Updates, and Settings.",
            "Use one safe-area strategy in the shell; pages do not add duplicate bottom insets.",
            "The page outlet reserves only the bottom-navigation height defined by the shell.",
            "Respect prefers-reduced-motion and keep keyboard focus visible.",
        ],
    )
    add_heading(doc, "11.2 Suggested Token Groups", 2)
    token_rows = [
        ("Color", "--color-brand, --color-surface, --color-text, --color-muted, --color-success, --color-warning"),
        ("Spacing", "--space-1 through --space-8, plus --page-padding"),
        ("Radius", "--radius-sm, --radius-md, --radius-lg, --radius-pill"),
        ("Typography", "--text-caption, --text-body, --text-title, --text-price"),
        ("Layout", "--content-max, --nav-height, --safe-bottom"),
        ("Motion", "--duration-fast, --duration-normal, --ease-standard"),
    ]
    add_table(doc, ["Token group", "Examples"], token_rows, [2000, 7360], size=9.1)

    add_heading(doc, "12. Models, Repositories, Services, and Storage", 1)
    add_heading(doc, "12.1 Model Responsibilities", 2)
    model_rows = [
        ("product.js", "Normalize uppercase item name, optional barcode/prices, active state, timestamps, and store/category IDs."),
        ("category.js", "Trim and uppercase category name; preserve store ownership and display order."),
        ("price-update.js", "Represent nullable previous prices, new prices, updater, device, and time."),
        ("profile.js", "Normalize display name, required email, optional mobile number, and photo URL."),
        ("store.js", "Represent hidden UUID, generated human-readable store code, name, and location."),
        ("device.js", "Represent store/user ownership, nickname, active state, and activity timestamps."),
    ]
    add_table(doc, ["Model helper", "Responsibility"], model_rows, [2300, 7060], size=9)

    add_heading(doc, "12.2 Repository Contracts", 2)
    repository_rows = [
        ("productRepository", "list, search, getById, getByBarcode, create, update, archive"),
        ("categoryRepository", "list, create, rename, reorder, archive"),
        ("priceUpdateRepository", "listRecent, listByProduct, record"),
        ("profileRepository", "getCurrent, update"),
        ("storeRepository", "getActive, create, findByStoreCode"),
        ("membershipRepository", "requestJoin, listMembers, approve, remove"),
        ("deviceRepository", "register, listActive, touchActivity, deactivate"),
    ]
    add_table(doc, ["Repository", "Stable operations"], repository_rows, [2600, 6760], size=9.2)
    add_callout(
        doc,
        "Why repositories matter",
        "Home and Browse should not care whether products come from sample arrays, a local cache, or Supabase. Only the repository implementation changes when the production data source changes.",
        fill=PALE_BLUE,
        accent=BLUE,
    )

    add_heading(doc, "12.3 Service Boundaries", 2)
    service_rows = [
        ("barcode-service.js", "Open scanner, return scanned code, and report permission/camera errors."),
        ("csv-service.js", "Parse, normalize, validate, preview, and export agreed product columns."),
        ("sync-service.js", "Coordinate cache, pending edits, retries, server updates, and conflict status."),
        ("profile-service.js", "Validate/prepare profile photos and profile-specific operations."),
        ("backup-service.js", "Build portable CSV backup and trigger download/share behavior."),
        ("auth-service.js", "Phase 8 wrapper for Supabase Auth session, email/password registration, login, logout, and recovery."),
    ]
    add_table(doc, ["Service", "Owns"], service_rows, [2500, 6860], size=9)

    add_heading(doc, "12.4 Storage Adapter", 2)
    add_body(
        doc,
        "The storage adapter is the only module that reads or writes localStorage. It uses versioned keys so cached data can be migrated safely.",
    )
    add_code_block(
        doc,
        """sarisync:v1:session-preferences
sarisync:v1:active-store
sarisync:v1:device
sarisync:v1:products-cache
sarisync:v1:categories-cache
sarisync:v1:price-updates-cache
sarisync:v1:pending-edits""",
        "Suggested key namespace",
    )

    doc.add_page_break()
    add_heading(doc, "13. Data, Authentication, and Security Boundaries", 1)
    add_heading(doc, "13.1 Prototype vs. Production", 2)
    production_rows = [
        ("Products/categories", "services/data.js arrays", "Repositories backed by Supabase and local cache."),
        ("Session", "Prototype localStorage object plus partial Supabase calls", "Supabase Auth session managed through auth-service."),
        ("Store joining", "Prototype form flow", "Generated store code, membership request, and owner approval."),
        ("PIN", "Still referenced by current auth prototype", "Removed from Version 1 architecture; email/password is the approved method."),
        ("Price history", "Derived from product sample fields", "Separate price_updates records from the Phase 6 model."),
        ("Device attribution", "Nickname stored in prototype session", "Registered device record linked to store and user."),
        ("Offline behavior", "Limited local values", "Versioned cache and pending queue; conflict policy remains later work."),
    ]
    add_table(doc, ["Concern", "Current prototype", "Approved direction"], production_rows, [2200, 3300, 3860], size=8.4)

    add_heading(doc, "13.2 Security Rules for Architecture", 2)
    add_list(
        doc,
        [
            "Never place a Supabase service-role or secret key in browser code.",
            "A publishable browser key does not replace Row Level Security.",
            "Every repository request must be scoped to the authenticated user's active store membership.",
            "The store code identifies a store; it is not an authentication secret.",
            "Email and password authenticate the person; active membership authorizes store access.",
            "Pages must not construct unrestricted table queries directly.",
            "Repository errors must be translated into user-safe messages without exposing credentials or internal details.",
        ],
    )
    add_callout(
        doc,
        "Phase boundary",
        "Phase 5 defines where authentication and authorization code belongs. Phase 8 implements authentication; the later security phase implements reviewed RLS policies.",
        fill=PALE_GOLD,
        accent="7A5A00",
    )

    add_heading(doc, "14. Safe Refactoring Sequence", 1)
    add_body(
        doc,
        "Refactor in small vertical slices. Each step must preserve the working UI and be verified before the next step starts.",
    )
    refactor_rows = [
        ("1", "Freeze behavior", "Record current screens, links, search, scanner, forms, and responsive checks.", "No intentional UI or behavior changes."),
        ("2", "Create src entry structure", "Add folders and entry files while existing files continue to run.", "App still opens from the current URL."),
        ("3", "Extract navigation", "Move active-tab, secondary-view, and Back-stack logic into navigation.js.", "All tabs and Back paths still work."),
        ("4", "Introduce central state", "Move session/store/device/navigation state behind state/store.js.", "Pages no longer read shared mutable variables directly."),
        ("5", "Extract page modules", "Move Browse, Updates, Settings, Profile, and product workflows one at a time.", "Each extracted page renders identically."),
        ("6", "Extract components", "Move product card, page header, toast, barcode button, and settings tile.", "Components receive data/callbacks only."),
        ("7", "Introduce models", "Centralize normalization and validation, including uppercase names and nullable optional prices.", "Forms and imports produce consistent records."),
        ("8", "Introduce repositories", "Wrap current arrays/local data behind stable interfaces.", "Pages no longer touch services/data.js directly."),
        ("9", "Split CSS", "Move tokens, base, layout, components, and page exceptions without selector changes first.", "Visual regression check passes."),
        ("10", "Add storage adapter", "Replace direct localStorage access with versioned methods.", "Session/cache behavior remains intact."),
        ("11", "Prepare production adapters", "Add Supabase-backed repositories after Phase 7 schema approval.", "Prototype adapter remains available for controlled testing."),
        ("12", "Remove compatibility code", "Delete old redirects/globals only after all references and tests are updated.", "No broken bookmarks or script paths."),
    ]
    add_table(doc, ["Step", "Change", "Action", "Exit check"], refactor_rows, [700, 1800, 3950, 2910], size=7.8, center_columns={0})

    add_heading(doc, "14.1 Stop Conditions", 2)
    add_list(
        doc,
        [
            "Stop if a move changes the UI, navigation history, or product behavior unexpectedly.",
            "Stop if the bottom navigation is duplicated or a page creates a second application shell.",
            "Stop if optional blank values become zero, undefined, or invented data.",
            "Stop if page code begins querying Supabase directly instead of using a repository.",
            "Stop if old and new modules both own the same state mutation.",
        ],
    )

    add_heading(doc, "15. Testing and Quality Gates", 1)
    test_rows = [
        ("Startup", "Unauthenticated user reaches Login; authenticated user reaches the intended initial view."),
        ("Navigation", "Tabs, deep entries, Back stack, Home avatar, and Settings Profile entry preserve history."),
        ("Home", "Search, scanner, Price Pulse, Quick Actions, and category preview behave unchanged."),
        ("Products", "Add, update, details, nullable wholesale price, barcode scan, and archive behavior remain correct."),
        ("Browse", "Search, category filtering, editing entry, and empty states work."),
        ("Updates", "Newest records appear first; missing previous price does not display as 0.00."),
        ("Settings", "Import/export remain administrative and navigation stays in one shell."),
        ("Responsive", "Small phone, large phone, and tablet widths share padding and avoid overlap."),
        ("Accessibility", "Keyboard focus, labels, live regions, contrast, and reduced motion remain usable."),
        ("Architecture", "Pages do not access localStorage or Supabase directly; components do not own business logic."),
        ("Data portability", "CSV preserves blank optional fields and the agreed column order."),
        ("Console", "No uncaught errors, missing scripts, broken imports, or invalid relative paths."),
    ]
    add_table(doc, ["Area", "Acceptance check"], test_rows, [2000, 7360], size=8.9)

    add_heading(doc, "16. Phase 5 Completion Checklist", 1)
    checklist_rows = [
        ("Current source structure inspected", "Completed"),
        ("Shared shell and navigation ownership defined", "Completed"),
        ("Target folder tree approved", "Pending review"),
        ("Page and component boundaries approved", "Pending review"),
        ("State shape and repository boundaries approved", "Pending review"),
        ("CSS layering and responsive rules approved", "Pending review"),
        ("Prototype and production data boundaries documented", "Completed"),
        ("Safe refactoring sequence approved", "Pending review"),
        ("Testing and acceptance gates approved", "Pending review"),
        ("Document approved for Phase 5 implementation", "Pending"),
    ]
    add_table(doc, ["Completion item", "Status"], checklist_rows, [7200, 2160], header_fill=PALE_GREEN, size=9.4, center_columns={1})

    add_heading(doc, "17. Approval and Handoff Record", 1)
    approval_rows = [
        ("Document version", "Draft 1.0"),
        ("Phase", "Phase 5 - Web Architecture"),
        ("Architecture approved?", "No - review required"),
        ("Implementation authorized?", "No - approve the staged plan first"),
        ("Reviewed by", ""),
        ("Review date", ""),
        ("Important changes requested", ""),
        ("Next phase after approval", "Complete safe Phase 5 refactoring, then begin Phase 7 Database Design"),
    ]
    add_table(doc, ["Item", "Your input"], approval_rows, [2900, 6460], header_fill=PALE_BLUE, size=9.5)

    add_callout(
        doc,
        "Recommended handoff",
        "Approve this document, create a separate implementation checklist from Section 14, and refactor one responsibility at a time. Do not combine the architecture refactor with database migration or major visual redesign.",
        fill=PALE_ORANGE,
    )

    doc.add_page_break()
    add_heading(doc, "Appendix A. Beginner-Friendly Glossary", 1)
    glossary = [
        ("Application shell", "The shared outer structure that stays mounted: page outlet, bottom navigation, and shared overlays."),
        ("Page module", "JavaScript responsible for rendering and coordinating one screen or workflow."),
        ("Component", "A reusable interface piece that receives data and callbacks."),
        ("State", "Information currently held by the app, such as the active store, products, search term, and sync status."),
        ("Central store", "One module that owns shared client state and notifies interested pages when it changes."),
        ("Model", "A helper that defines, normalizes, validates, and serializes one kind of record."),
        ("Service", "Logic that performs an operation, such as barcode scanning or CSV parsing."),
        ("Repository", "A stable data interface that hides whether information comes from mock data, cache, or Supabase."),
        ("Storage adapter", "The only module allowed to access browser persistence such as localStorage."),
        ("Dependency wiring", "Connecting pages to the specific state, repositories, and services they need when the app starts."),
        ("Back stack", "An ordered list of previously visited views used to return to the correct screen."),
        ("Normalization", "Turning inputs into one consistent form, such as trimming and uppercasing category names."),
        ("Validation", "Checking that data follows rules before it is saved."),
        ("Data contract", "The agreed fields, types, optional values, and behavior of a record."),
        ("Adapter", "A module that makes one storage system follow the repository interface."),
        ("RLS", "Row Level Security: database rules that restrict which store rows an authenticated user may access."),
        ("Regression", "A previously working behavior that breaks after a change."),
    ]
    add_table(doc, ["Term", "Simple meaning"], glossary, [2100, 7260], size=9.2)

    add_heading(doc, "Appendix B. Evidence Reviewed", 1)
    evidence = [
        ("ARCHITECTURE.md", "Current documented folder responsibilities and shared secondary-page approach."),
        ("index.html", "Shared shell, Home content, bottom navigation, shared sheet, and toast roots."),
        ("pages/home/home.js", "Current centralized rendering, navigation, forms, CSV, and product mutations."),
        ("pages/auth/*", "Current authentication/store setup page and partial Supabase connection."),
        ("pages/categories/*", "Reusable standalone category view and localStorage transfer."),
        ("services/*.js", "Prototype data, formatting, responsive tokens, barcode scanning, and Supabase client."),
        ("SariSync_V1_Feature_Roles_Requirements.docx", "Version 1 scope, roles, limits, and requirements."),
        ("SariSync_Phase_3_Screen_Inventory.docx", "View purposes, entry/exit behavior, and navigation expectations."),
        ("SariSync_Phase_6_Data_Model.docx", "Approved entity, field, relationship, and business-rule decisions."),
    ]
    add_table(doc, ["Source", "How it informed Phase 5"], evidence, [3400, 5960], size=8.8)
    add_callout(
        doc,
        "Accuracy note",
        "This document describes the source state inspected on July 29, 2026. Future code changes should update ARCHITECTURE.md and this document's version history.",
        fill=PALE_GOLD,
        accent="7A5A00",
    )

    doc.core_properties.title = "SariSync Phase 5 Web Architecture"
    doc.core_properties.subject = "Version 1 browser architecture, modular structure, and staged refactoring plan"
    doc.core_properties.author = "SariSync Project"
    doc.core_properties.keywords = "SariSync, Phase 5, web architecture, HTML, CSS, JavaScript, repository, state"
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build_document()
